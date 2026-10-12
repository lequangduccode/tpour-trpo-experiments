"""Create the Vietnamese Word report and editable PowerPoint from collected evidence."""
import json
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
ASSETS = OUT / "assets"
EV = OUT / "evidence"
REPO = "https://github.com/lequangduccode/tpour-trpo-experiments"
PAPER = "https://arxiv.org/abs/2606.17664"


def j(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def architecture():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    fig, ax = plt.subplots(figsize=(12, 5.4))
    fig.patch.set_facecolor("white")
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5.4)
    ax.axis("off")
    def box(x, y, w, h, txt, color="#edf3ff"):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=.08", ec="#aac0de", fc=color))
        ax.text(x+w/2, y+h/2, txt, ha="center", va="center", fontsize=12, color="#14243d")
    def arrow(a, b, label=None):
        ax.annotate("", b, a, arrowprops=dict(arrowstyle="->", color="#3767e8", lw=2))
        if label:
            ax.text((a[0]+b[0])/2, (a[1]+b[1])/2 + .12, label, fontsize=10, ha="center")
    box(.2, 2.2, 2.1, 1.2, "Q / D đúng mốc\nD lệch mốc")
    box(3, 3.5, 2.7, 1.1, "Encoder hiện tại\nθ · nhận gradient")
    box(3, 1.1, 2.7, 1.1, "Encoder tham chiếu\nθ_ref · no_grad", "#e6f7f2")
    box(6.4, 3.5, 2.3, 1.1, "Điểm tương đồng\nCE + TRPO")
    box(6.4, 1.1, 2.3, 1.1, "Queue MoCo\nEmbeddings cũ", "#e6f7f2")
    box(9.5, 3.5, 2.2, 1.1, "Backpropagation\nAdamW")
    box(9.5, 1.1, 2.2, 1.1, "Momentum\nm = 0.9999", "#e6f7f2")
    arrow((2.3, 3), (3, 4))
    arrow((2.3, 2.6), (3, 1.65))
    arrow((5.7, 4.05), (6.4, 4.05))
    arrow((5.7, 1.65), (6.4, 1.65))
    arrow((5.7, 2), (6.5, 3.55))
    arrow((7.55, 2.2), (7.55, 3.5))
    arrow((8.7, 4.05), (9.5, 4.05))
    arrow((10.6, 3.5), (10.6, 2.2))
    arrow((9.5, 1.45), (5.7, 1.45))
    ax.text(6, .3, "Sơ đồ tự vẽ từ phương trình 3–5 và triển khai tpour.py; không phải ảnh giao diện demo.",
            ha="center", fontsize=10, color="#64748b")
    fig.savefig(ASSETS / "architecture.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tcpr.append(shade)


class Report:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
        sec.left_margin, sec.right_margin = Cm(2.25), Cm(2)
        sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
        style = self.doc.styles["Normal"]
        style.font.name, style.font.size = "Times New Roman", Pt(11.5)
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(7)
        for s in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            style = self.doc.styles[s]
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor.from_string("163B68")
        self.doc.styles["Heading 1"].font.size = Pt(18)
        self.doc.styles["Heading 2"].font.size = Pt(13)
        header = sec.header.paragraphs[0]
        header.text = "TPOUR / TRPO  |  BÁO CÁO PHÂN TÍCH VÀ THỬ NGHIỆM"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string("64748B")
        footer = sec.footer.paragraphs[0]
        footer.alignment = 2
        run = footer.add_run("12/10/2026  •  Trang ")
        run.font.size = Pt(9)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)
        self.doc.core_properties.title = "TPOUR/TRPO — Báo cáo phân tích và thử nghiệm"
        self.doc.core_properties.subject = "Temporal Preference Optimization for Unsupervised Retrieval"
        self.doc.core_properties.author = ""
        self.doc.core_properties.keywords = "TPOUR, TRPO, Contriever, temporal retrieval, CPU demo"

    def page(self, heading):
        self.doc.add_page_break()
        self.doc.add_heading(heading, 1)

    def p(self, text, bold=False):
        p = self.doc.add_paragraph()
        p.alignment = 3
        p.add_run(text).bold = bold
        return p

    def h(self, text):
        self.doc.add_heading(text, 2)

    def bullets(self, lines):
        for line in lines:
            self.doc.add_paragraph(line, style="List Bullet")

    def table(self, heads, rows):
        table = self.doc.add_table(rows=1, cols=len(heads))
        table.style = "Table Grid"
        for cell, txt in zip(table.rows[0].cells, heads):
            cell.text = str(txt)
            cell_shading(cell, "163B68")
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, txt in zip(cells, row):
                cell.text = str(txt)
                if i % 2 == 0:
                    cell_shading(cell, "F0F4FA")
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    for run in p.runs:
                        run.font.size = Pt(10)
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
        self.doc.add_paragraph()

    def image(self, name, caption, width=6.4):
        p = self.doc.add_paragraph()
        p.alignment = 1
        p.add_run().add_picture(str(ASSETS / name), width=Inches(width))
        p = self.doc.add_paragraph(caption)
        p.alignment = 1
        for r in p.runs:
            r.italic = True
            r.font.size = Pt(9)

    def code(self, text):
        p = self.doc.add_paragraph()
        for i, line in enumerate(text.splitlines()):
            if i:
                p.add_run().add_break()
            r = p.add_run(line)
            r.font.name, r.font.size = "Consolas", Pt(8.5)
        p.paragraph_format.space_after = Pt(10)


def build_word(summary, meta, ranking):
    r = Report()
    doc = r.doc
    doc.add_paragraph("BÁO CÁO NGHIÊN CỨU VÀ THỰC HÀNH", style="Subtitle")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(50)
    p.add_run("TPOUR / TRPO").font.size = Pt(32)
    doc.add_heading("Tối ưu ưu tiên thời gian cho hệ thống truy xuất không giám sát", 1)
    r.p("Phân tích thuật toán • Triển khai PyTorch • Thực nghiệm CPU • Bằng chứng demo", True)
    r.image("architecture.png", "Luồng học kết hợp ngữ nghĩa và tín hiệu thời gian", width=6.1)
    r.p("Người thực hiện: ...............................................................")
    r.p("Lớp / mã sinh viên: ...........................................................")
    r.p("Giảng viên hướng dẫn: ........................................................")
    r.p("Ngày hoàn thiện: 12/10/2026")
    r.p("Bài báo: Temporal Preference Optimization for Unsupervised Retrieval — HyunJin Kim, Jaejun Shim, Young Jin Kim, JinYeong Bak.")
    r.p("Phạm vi: bản triển khai gọn và kiểm tra khả năng chạy; chưa tái lập đầy đủ kết quả công bố.", True)
    r.page("Tóm tắt và phạm vi báo cáo")
    r.p("Báo cáo nghiên cứu cách bổ sung tín hiệu thời gian vào dense retrieval. Một truy vấn có thể đúng về chủ đề nhưng vẫn cần tài liệu thuộc đúng giai đoạn. TPOUR kết hợp học tương phản với Temporal Retrieval Preference Optimization (TRPO), trong đó độ chênh điểm giữa tài liệu đúng mốc và lệch mốc được so với encoder tham chiếu cập nhật bằng momentum [1, 2].")
    r.p("Phần thực hành xây dựng bộ mã một thiết bị, hỗ trợ encoder nhỏ để demo và facebook/contriever để kiểm tra mô hình thật. Bộ mã có chuẩn bị dữ liệu, huấn luyện, lưu/nạp checkpoint, retrieval, đánh giá chỉ số, nội suy tham số và script BEIR. Các kết quả được phân loại theo nguồn nhằm tránh nhầm giữa số liệu của tác giả và số liệu tự đo.")
    r.table(["Nhóm bằng chứng", "Đã thực hiện", "Kết luận phù hợp"], [
        ["Demo giả lập", "3 seed; 150 bước; 32 cặp train; 8 query test", "Pipeline chạy; bài toán dễ và có hiệu ứng trần"],
        ["Kiểm tra tự động", "5 kiểm tra thành công", "Một số bất biến và chiều gradient đúng"],
        ["Contriever thật", "2 bước trên 8 mẫu, ngày 10/10/2026", "Mô hình thật forward/backward và lưu/nạp được"],
        ["BEIR / temporal QA", "Chưa chạy benchmark đầy đủ", "Chưa đủ bằng chứng tái lập hoặc so sánh SOTA"]])
    r.p("Kết quả seed 42: nDCG@5 của encoder chưa huấn luyện là 0,3028; chỉ CE và CE + TRPO đều đạt 1,0000. Do đối chứng cũng đạt tối đa, kết quả này không chứng minh lợi thế riêng của TRPO. Hai seed bổ sung chỉ kiểm tra độ nhạy của demo đối với khởi tạo, không thay thế đánh giá trên dữ liệu thực.")
    r.h("Cách đọc báo cáo")
    r.p("[1] là bài báo; [2] là code tác giả; [3]–[5] là tài liệu mô hình và benchmark. Mục 1–5 trình bày cơ sở và thiết kế; mục 6–10 trình bày thực nghiệm và bằng chứng; mục 11–14 phân tích giới hạn, cách chạy lại và hướng tiếp theo. Các bảng kết quả nội bộ đều ghi rõ dữ liệu giả lập hoặc smoke test.")
    r.page("Nội dung và thuật ngữ")
    contents = ["1. Bài toán truy xuất có yếu tố thời gian", "2. Cơ sở: Contriever, MoCo và preference learning", "3. Hàm mất mát TRPO và cập nhật tham chiếu", "4. Kiến trúc bản triển khai", "5. Dữ liệu và kiểm soát rò rỉ", "6. Thiết kế và cấu hình thực nghiệm", "7. Kết quả retrieval và thử nhiều seed", "8. Phân tích loss và kiểm tra tính đúng", "9. Kiểm tra với Contriever thật", "10. Ảnh chụp demo và cách diễn giải", "11. Nội suy tham số theo thời gian", "12. Phân tích giới hạn và khác biệt với bài gốc", "13. Hướng dẫn chạy lại và trình diễn", "14. Kế hoạch đánh giá mở rộng và kết luận", "Tài liệu tham khảo và hỏi đáp thuyết trình"]
    for line in contents:
        doc.add_paragraph(line)
    r.table(["Thuật ngữ", "Ý nghĩa trong báo cáo"], [["Dense retriever", "Mã hóa query và document thành vector rồi xếp hạng bằng độ tương đồng"], ["Temporal alignment", "Phù hợp ngữ cảnh thời gian của truy vấn"], ["Momentum encoder", "Encoder tham chiếu cập nhật chậm từ trọng số encoder hiện tại"], ["Queue", "Bộ đệm embeddings làm negative cho các batch sau"], ["Qrels", "Nhãn độ liên quan dùng để đánh giá retrieval"], ["Smoke test", "Kiểm tra đường chạy cơ bản; không đánh giá đầy đủ chất lượng"]])
    r.page("1. Bài toán truy xuất có yếu tố thời gian")
    r.h("1.1. Vì sao tương đồng ngữ nghĩa chưa đủ?")
    r.p("Một kho tài liệu nhiều năm có thể chứa nhiều phiên bản của cùng một chủ đề. Khi hỏi ‘Tổ chức hiện cung cấp dịch vụ thế nào?’, văn bản cũ và mới cùng nhắc đến tổ chức đó nên đều có điểm tương đồng cao. Tuy nhiên, văn bản cũ có thể mô tả quy trình đã thay đổi. Vấn đề là xếp hạng đúng chủ đề và đúng ngữ cảnh thời gian, thay vì chỉ gần nhau về từ ngữ [1].")
    r.table(["Truy vấn / tài liệu", "Ví dụ giả định", "Vai trò"], [["Q", "How does Organization32 currently provide its service?", "Truy vấn có thời gian ngầm định"], ["D đúng mốc", "Uses a digital portal and provides online mobile service.", "Phiên bản 2021 được ưu tiên"], ["D lệch mốc", "Uses paper forms and provides office desk service.", "Phiên bản 2018, cùng chủ đề"]])
    r.h("1.2. Thời gian tường minh và ngầm định")
    r.p("Thời gian tường minh xuất hiện trực tiếp như ‘năm 2019’. Thời gian ngầm định dùng ‘hiện tại’, ‘năm nay’ hoặc phụ thuộc bối cảnh người hỏi. Trong TPOUR, encoder chuyên biệt theo một mốc được huấn luyện để ưu tiên corpus của mốc đó. Một encoder căn chỉnh 2018 không tự hiểu ‘hiện tại’ là ngày chạy máy; nó phản ánh mốc học. Đây là điểm cần giải thích khi trình bày [1].")
    r.h("1.3. Mục tiêu của phần thực hành")
    r.bullets(["Kiểm tra hàm loss và cơ chế tham chiếu/queue có hoạt động trong PyTorch.", "Tạo đối chứng chỉ CE để không gán mọi cải thiện cho TRPO.", "Theo dõi retrieval, loss, lưu/nạp mô hình và khả năng chạy mô hình thật.", "Xác định phần nào đã kiểm chứng và phần nào cần dữ liệu/GPU để đánh giá tiếp."])
    r.p("Bản demo không xây dựng một bộ tìm kiếm tổng quát theo ngày tháng. Nó chọn một mốc đích cho dữ liệu giả lập và đo thứ hạng tài liệu đúng phiên bản.")
    r.page("2. Cơ sở: Contriever, MoCo và ưu tiên")
    r.h("2.1. Dense retrieval và mean pooling")
    r.p("Encoder biến query Q và tài liệu D thành vector q và d. Với Contriever, bản thực hành sử dụng AutoModel và mean pooling có attention mask: cộng embeddings của token hợp lệ rồi chia cho số token. Điểm retrieval là tích vô hướng qᵀd. Nếu bật chuẩn hóa L2, tích vô hướng trở thành cosine; mặc định nhánh Contriever trong thực hành không chuẩn hóa [3].")
    r.code("q = encoder(Q); d = encoder(D)\nscore(Q, D) = dot(q, d)\nmean_pool = sum(hidden * attention_mask) / sum(attention_mask)")
    r.h("2.2. Học tương phản và MoCo")
    r.p("Học tương phản tăng điểm của cặp liên quan và giảm tương quan với negative. MoCo duy trì một encoder tham chiếu cập nhật chậm cùng hàng đợi vector tài liệu. Queue tái sử dụng negative từ batch trước, làm tăng số negative mà không cần batch lớn đồng thời. Queue lớn hơn vẫn tốn RAM/VRAM và phép nhân ma trận; cơ chế này không làm toàn bộ huấn luyện trở nên miễn phí [1, 2].")
    r.h("2.3. Ý tưởng ưu tiên từ DPO")
    r.p("DPO so sánh phản hồi ưu tiên với phản hồi bị từ chối tương đối so với chính sách tham chiếu. TRPO trong bài này thay điểm log-likelihood bằng độ tương đồng retrieval, và ưu tiên đến từ khác biệt giữa các corpus theo thời gian. Không cần reward model riêng; vẫn cần cặp tài liệu cùng chủ đề với tín hiệu thời gian có ý nghĩa [1, 5].")
    r.p("TRPO ở đây viết tắt của Temporal Retrieval Preference Optimization. Không nhầm với Trust Region Policy Optimization trong reinforcement learning.", True)
    r.h("2.4. Hai loại encoder trong repository")
    r.p("Encoder toy của demo là bảng embedding 4096 × 64, ánh xạ từ bằng SHA-256, lấy trung bình rồi chuẩn hóa. Nó được khởi tạo ngẫu nhiên và không có kiến thức ngôn ngữ tiền huấn luyện. Encoder Contriever thật dùng trọng số facebook/contriever. Vì thế ‘untrained’ trong bảng demo không phải baseline Contriever của bài báo.")
    r.page("3. TRPO: công thức và ý nghĩa")
    r.h("3.1. Chênh lệch điểm ưu tiên")
    r.code("S_theta(Q,D) = dot(enc_theta(Q), enc_theta(D)) / T\nDelta_theta = S_theta(Q,D_preferred) - S_theta(Q,D_rejected)\nDelta_ref   = S_ref(Q,D_preferred) - S_ref(Q,D_rejected)\nL_TRPO = -mean(logsigmoid(beta * (Delta_theta - Delta_ref)))")
    r.p("Delta_theta thể hiện mức encoder hiện tại thích tài liệu đúng mốc hơn tài liệu lệch mốc. Trừ Delta_ref đặt mức ưu tiên này tương đối so với encoder tham chiếu. Khi chênh lệch hiện tại tốt hơn tham chiếu, đối số sigmoid tăng và loss giảm. Nếu bỏ toàn bộ phần tham chiếu, đó là pairwise ranking loss khác với công thức TRPO của bài [1].")
    r.h("3.2. Loss tương phản và loss tổng")
    r.code("L_CE = -log(exp(S_theta(Q,D_preferred)) /\n        (exp(S_theta(Q,D_preferred)) + sum(exp(S(Q,k_neg)))))\nL_total = lambda * L_CE + (1 - lambda) * L_TRPO")
    r.p("Trong bản gọn, positive dùng encoder hiện tại; negative là vector reference đã lưu trong queue. Tất cả logits CE được chia temperature T. Mặc định λ = 0,925 nên trọng số TRPO là 0,075. β = 0,3 điều khiển độ nhạy ưu tiên. Đây là hệ số loss, không phải tỉ lệ dữ liệu hoặc phần trăm tài liệu đúng mốc. Ảnh hưởng gradient còn phụ thuộc thang điểm, T và kích thước queue.")
    r.h("3.3. Ổn định số và gradient")
    r.p("Code dùng log-sigmoid để tránh tính sigmoid rồi log một cách kém ổn định. Điểm reference được detach và mô hình reference có requires_grad=False. Khi Delta_theta = Delta_ref, TRPO loss bằng ln(2), khoảng 0,6931; gradient vẫn đẩy điểm preferred lên và rejected xuống. Kiểm tra tự động xác nhận chiều này trong trường hợp tensor đơn giản.")
    r.p("Bản code tác giả và công thức lý thuyết không hoàn toàn trùng chi tiết CE. Báo cáo trình bày riêng các khác biệt ở mục 12 thay vì nhận định đây là tái lập bit-for-bit.")
    r.page("4. Kiến trúc và quy trình cập nhật")
    r.image("architecture.png", "Hình 1. Luồng encoder hiện tại, reference, loss và queue của bản triển khai.")
    r.h("4.1. Một bước huấn luyện")
    r.bullets(["Lấy batch query/preferred/rejected; tùy chọn xóa ngẫu nhiên từ.", "Encoder hiện tại mã hóa ba nhóm văn bản, reference mã hóa dưới no_grad và eval mode.", "Tính CE với queue cũ, tính TRPO theo chênh lệch tương đối; cộng loss theo λ.", "Backpropagate, clip gradient norm ở 1,0, AdamW cập nhật encoder hiện tại.", "Cập nhật reference bằng momentum, sau đó enqueue embeddings reference vừa tính."])
    r.code("theta_ref <- m * theta_ref + (1 - m) * theta\nm = 0.9999")
    r.p("m lớn giúp reference thay đổi chậm. Reference không phải mô hình nền đóng băng suốt quá trình: trọng số vẫn thay đổi bằng momentum. Queue khởi động bằng embeddings của dữ liệu train thật để bước đầu có negative. Queue vòng xử lý batch vượt cuối buffer bằng cách chia đoạn, không yêu cầu queue size chia hết batch size.")
    r.h("4.2. Truy xuất sau huấn luyện")
    r.p("Chỉ encoder hiện tại cần thiết cho inference. Tài liệu được mã hóa trước; query được mã hóa rồi tính điểm với toàn corpus để lấy top-k. Dùng encoder nào thì index phải được mã hóa bằng encoder đó. Khi thay checkpoint hoặc nội suy trọng số, phải mã hóa lại corpus, không tái sử dụng index của encoder khác.")
    r.page("5. Dữ liệu và kiểm soát rò rỉ")
    r.h("5.1. Dữ liệu giả lập dùng cho demo")
    r.p("Có 40 tổ chức giả định với hai phiên bản dịch vụ: paper/office ở 2018 và digital/online ở 2021. Tổ chức 0–31 tạo 32 bộ ba huấn luyện. Tổ chức 32–39 tạo 8 query và 16 tài liệu test. Với mỗi query test, tài liệu phiên bản mới của đúng tổ chức có qrel = 1; mọi tài liệu khác có qrel = 0. Nhãn và mốc của demo được dựng theo kịch bản, không phải dữ liệu thu thập ngoài đời.")
    r.table(["Phần", "Số lượng", "Vai trò / lưu ý"], [["Train", "32 bộ ba", "Q + preferred 2021 + rejected 2018"], ["Test query", "8", "Thực thể không trùng train"], ["Test corpus", "16", "8 thực thể × 2 phiên bản"], ["Qrels", "8 dòng positive", "Chỉ dùng cho đánh giá"], ["Mẫu tác giả", "3.000 dòng upstream", "Smoke test lấy 8 dòng đầu"]])
    r.h("5.2. Dữ liệu thực theo snapshot")
    r.p("Theo bài báo, corpus thời gian cần giữ tài liệu mới hoặc có nội dung thay đổi và loại tài liệu gold của tập test [1]. Timestamp là mốc thu thập corpus, không đồng nhất với mọi năm xuất hiện trong văn bản. Tài liệu đúng mốc và lệch mốc cần tương đồng chủ đề; lấy negative hoàn toàn khác chủ đề sẽ làm bài toán quá dễ và ít kiểm tra được thời gian.")
    r.h("5.3. Bộ chuẩn bị dữ liệu của bản gọn")
    r.p("Lệnh prepare chia đôi tài liệu tại mốc đích để tạo pseudo-query và positive, rồi tìm tài liệu khác mốc có điểm cosine cao nhất làm rejected. Đây là heuristic thực hành, khác pipeline mining của tác giả. Tài liệu ngắn hơn 20 từ bị bỏ; phải còn ít nhất hai bộ ba hợp lệ. Bộ công cụ giữ vector trong RAM nên phù hợp tập nhỏ.")
    r.p("Tách theo thực thể giảm một dạng rò rỉ, nhưng mẫu câu chung và từ digital/online vẫn giúp mô hình đoán dễ. Với dữ liệu thật cần tách theo tài liệu, đối chiếu phiên bản, loại near-duplicate và kiểm tra cùng nội dung dưới ID khác. Tùy chọn exclude chỉ loại ID trùng, không tự phát hiện mọi rò rỉ.")
    r.page("6. Thiết kế và cấu hình thực nghiệm")
    r.h("6.1. Câu hỏi kiểm chứng")
    r.p("RQ-A: pipeline tính loss, cập nhật reference và queue có hoạt động? RQ-B: sau train có retrieval và nạp lại checkpoint được không? RQ-C: thêm TRPO có tạo lợi ích so với chỉ CE trên cùng demo? RQ-D: có chạy được mô hình Contriever thật trên môi trường hiện tại? Những câu hỏi này hẹp hơn các câu hỏi khoa học về tổng quát hóa thời gian của bài báo.")
    r.table(["Tham số", "Demo toy", "Contriever smoke test"], [["Bước / batch", "150 / 8", "2 / 2"], ["Số mẫu train", "32", "8 mẫu tác giả"], ["Learning rate", "0,01", "1e-6"], ["Queue / max length", "256 / 256", "8 / 64"], ["Encoder / dim", "Toy / 64", "facebook/contriever / 768"], ["Chuẩn hóa vector", "L2", "Không"], ["Seed", "42, 43, 44", "42"], ["λ / β / T / m", "0,925 / 0,3 / 0,05 / 0,9999", "Giống demo"], ["Xóa từ ngẫu nhiên", "0%", "10%"], ["AdamW", "β1=0,9; β2=0,98; ε=1e-6", "Giống demo"]])
    r.h("6.2. Đối chứng")
    r.p("Mỗi seed bắt đầu từ cùng encoder toy khởi tạo rồi sao chép cho hai nhánh. Nhánh chỉ CE có λ = 1; nhánh TPOUR có λ = 0,925. Các tham số còn lại và seed lấy mẫu được giữ tương ứng. ‘Untrained’ là encoder toy ban đầu. CE-only vẫn tính TRPO để ghi log nhưng TRPO không đóng góp vào loss tổng.")
    r.h("6.3. Môi trường và nguồn bằng chứng")
    r.p(f"Chạy bổ sung ngày 12/10/2026: Python {meta['python']}, PyTorch {meta['torch']}; CUDA available = {meta['cuda_available']}. Đầu ra lệnh thực tế, timestamp theo Asia/Bangkok, hash code và commit nền được lưu trong reports/evidence. Mỗi seed được chạy bằng tiến trình Python riêng.")
    r.page("7. Kết quả retrieval trên demo")
    r.table(["Mô hình (seed 42)", "nDCG@5", "Recall@5", "MRR@5", "nDCG@10"], [[label] + [f"{summary['42'][key][m]:.4f}" for m in ("nDCG@5", "Recall@5", "MRR@5", "nDCG@10")] for key, label in (("untrained", "Toy chưa train"), ("ce_only", "Chỉ CE"), ("tpour", "CE + TRPO"))])
    r.image("metrics.png", "Hình 2. Các điểm đo trên toàn bộ tập test giả lập; không phải điểm BEIR.")
    r.h("7.1. Cách hiểu chỉ số")
    r.p("nDCG@k ưu tiên đặt tài liệu liên quan ở đầu danh sách và chuẩn hóa theo thứ hạng lý tưởng. Recall@k trong code là số tài liệu liên quan tìm thấy chia tổng tài liệu liên quan. MRR@k lấy nghịch đảo hạng tài liệu liên quan đầu tiên, hoặc 0 nếu top-k không có. Demo có đúng một positive cho mỗi query, nên Recall@k tương đương hit-rate trong tập này; không đồng nhất hai khái niệm trên tập nhiều positive.")
    r.code("DCG@k = sum((2^rel_i - 1) / log2(i + 1)), i=1..k\nnDCG@k = DCG@k / IDCG@k\nRecall@k = |relevant ∩ top-k| / |relevant|\nMRR@k = mean(1 / rank_first_relevant), rank <= k")
    r.p("Cả hai mô hình đã train đạt tối đa, còn encoder khởi tạo ngẫu nhiên thấp hơn. Kết luận phù hợp là mô hình học được quy luật của demo. Chênh lệch giữa CE + TRPO và chỉ CE bằng 0 ở các chỉ số này nên chưa xác nhận giá trị gia tăng của TRPO.")
    r.page("7.2. Kiểm tra bổ sung với ba seed")
    r.table(["Seed", "Chưa train N@5", "Chỉ CE N@5", "CE+TRPO N@5", "Δ TRPO − CE"], [[s] + [f"{summary[s][k]['nDCG@5']:.4f}" for k in ("untrained", "ce_only", "tpour")] + [f"{summary[s]['tpour']['nDCG@5'] - summary[s]['ce_only']['nDCG@5']:+.4f}"] for s in ("42", "43", "44")])
    r.p("Ba seed 42, 43 và 44 là các lần chạy thật của cùng kịch bản giả lập, không phải dữ liệu lấy từ bài báo. Bảng giúp phân biệt biến động khởi tạo ngẫu nhiên với hiệu quả huấn luyện. Vì test chỉ có 8 query và dùng chung template, độ ổn định cao không chứng minh tổng quát hóa trên corpus mới.")
    r.h("7.3. Ví dụ thứ hạng từ checkpoint")
    selected = ranking["selected"]
    base_rank = next(i+1 for i,d in enumerate(ranking["rankings"]["base"][selected]) if d["relevant"])
    tp_rank = next(i+1 for i,d in enumerate(ranking["rankings"]["tpour"][selected]) if d["relevant"])
    r.p(f"Ví dụ được chọn để hiển thị: query {selected}, hỏi về Organization{selected}. Trong encoder toy ban đầu, tài liệu đúng đứng hạng {base_rank}/16; sau CE + TRPO đứng hạng {tp_rank}/16. Ví dụ này được chọn theo truy vấn có hạng baseline kém nhất trong 8 query, nên mang tính minh họa; bảng chỉ số phía trước vẫn tính trên toàn bộ test.")
    r.table(["Nhánh", "Hạng 1", "Điểm", "Đúng qrel?"], [[name, ranking["rankings"][key][selected][0]["id"], f"{ranking['rankings'][key][selected][0]['score']:.4f}", "Có" if ranking["rankings"][key][selected][0]["relevant"] else "Không"] for key,name in (("base","Chưa train"),("tpour","CE + TRPO"))])
    r.h("7.4. Những kết luận chưa được hỗ trợ")
    r.bullets(["Không suy ra mức cải thiện của Contriever từ điểm encoder toy.", "Không suy ra thời gian inference trên GPU hoặc khả năng chạy thiết lập A100 trên RTX 3090.", "Không suy ra khả năng xử lý tiếng Việt từ dữ liệu tiếng Anh giả lập.", "Không khẳng định học ưu tiên thời gian khi gặp thay đổi nội dung tinh tế hoặc query mơ hồ ngoài template.", "Không xem ba seed nhỏ là bằng chứng thống kê cho một benchmark thực."])
    r.page("8. Loss huấn luyện và kiểm tra tính đúng")
    r.image("loss.png", "Hình 3. Loss từ log thật của seed 42; trục dọc log giúp quan sát đoạn cuối.")
    ce_log = j(EV / "ce_only_train_log.json")
    tp_log = j(EV / "tpour_train_log.json")
    r.table(["Nhánh / bước 150", "Loss tổng", "CE", "TRPO ghi nhận"], [[name] + [f"{logs[-1][key]:.6f}" for key in ("loss","ce","trpo")] for name,logs in (("Chỉ CE",ce_log),("CE + TRPO",tp_log))])
    r.p("Loss giảm cho thấy tối ưu trên batch train tiến triển, nhưng không thay thế đánh giá retrieval. Hai loss tổng có trọng số khác nhau nên không so sánh số nhỏ hơn như một thước đo chất lượng độc lập. Nhánh CE-only cũng có TRPO loss giảm do ưu tiên ở demo trùng với quy luật CE dễ học; điều này củng cố nhu cầu đối chứng.")
    r.h("8.1. Năm kiểm tra tự động")
    r.table(["Kiểm tra", "Điều được xác nhận"], [["Loss và gradient", "Tại gap tương đương, loss ≈ ln(2); gradient đúng chiều preferred/rejected"], ["Reference detach", "Tensor reference không nhận gradient từ TRPO"], ["Queue và momentum", "Buffer wrap đúng; cập nhật trọng số đúng hệ số"], ["Metrics", "nDCG/Recall/MRR khớp ví dụ tính tay"], ["Input upstream", "Đọc đúng question và context dạng chuỗi / dict"]])
    r.p("Kết quả chạy lại: 5/5 đạt. Những kiểm tra này không chứng minh toàn bộ thuật toán hoặc dữ liệu hoàn hảo; chúng xác nhận các tình huống có đáp án mong đợi rõ ràng.")
    r.page("9. Kiểm tra với Contriever thật")
    r.h("9.1. Bằng chứng đã có từ ngày 10/10/2026")
    r.p("Tám dòng đầu của file 2018_sample_train_trimmed.jsonl được lưu tại data/official_smoke.jsonl. AutoTokenizer và AutoModel tải facebook/contriever. Chạy 2 bước trên CPU, batch 2, max length 64, queue 8. Đầu ra nằm trong results/contriever_smoke; checkpoint khoảng 438 MB nằm ở outputs/contriever_smoke/encoder.pt trên máy thực nghiệm.")
    real = j(ROOT / "results/contriever_smoke/train_log.json")
    r.table(["Bước", "Loss tổng", "CE", "TRPO"], [[v["step"]] + [f"{v[k]:.8f}" for k in ("loss","ce","trpo")] for v in real])
    r.p("Loss tổng bước 2 cao hơn bước 1 không phải bằng chứng huấn luyện hỏng: batch thay đổi, queue cập nhật và quá trình chỉ có hai điểm. Cũng không thể kết luận hội tụ từ hai bước. Mục đích là kiểm tra forward, backward, optimizer, momentum, queue và lưu mô hình thật.")
    r.h("9.2. Nạp lại checkpoint")
    r.p("Checkpoint đã được nạp lại để retrieval trên tập giả lập 8 query/16 tài liệu. Kết quả có sẵn trong synthetic_reload_check.json. Bộ test này không phải corpus huấn luyện tác giả hoặc BEIR; kết quả cao có thể do Contriever tiền huấn luyện đã hiểu quy luật văn bản. Do chưa có baseline Contriever cùng tập được đo trước fine-tune, không gán điểm này cho hiệu quả hai bước TRPO.")
    r.h("9.3. Khoảng cách tới thiết lập công bố")
    r.p("Thiết lập trong phần phụ lục bài báo mô tả fine-tune đầy đủ với A100, queue lớn, chuỗi dài và lịch học có warmup. Bản hiện tại dùng CPU, queue 8 và 2 bước. Để tái lập, cần chuẩn hóa phiên bản bài/code, dataset, seed, mining, optimizer/scheduler, hardware và evaluation protocol. Việc mô hình có khoảng 110 triệu tham số không tự bảo đảm thiết lập đầy đủ vừa mọi GPU lab [1, 2].")
    r.p("Các log Contriever được ghi nhận là bằng chứng lịch sử ngày 10/10/2026; ảnh demo chụp ngày 12/10/2026 hiển thị thông tin này rõ ràng.")
    r.page("10. Ảnh chụp demo: tổng quan")
    r.image("demo_01_overview.png", "Hình 4. Ảnh chụp trình duyệt Chrome từ reports/demo_viewer.html, chứa kết quả đo thật trên demo giả lập.", width=6.1)
    r.h("Cách trình bày hình trước giảng viên")
    r.p("Bắt đầu từ quy mô nhỏ: 32 cặp train, 8 query và 16 tài liệu test. Chỉ vào bảng ba nhánh để giải thích baseline và đối chứng. Nhấn mạnh dữ liệu giả lập và việc hai nhánh đã train cùng đạt 1,0. Không dùng ảnh này như bằng chứng hệ thống đạt hiệu quả tương tự bài báo.")
    r.p("Viewer được tạo từ JSON và log thực tế, có thể mở offline. Các điểm retrieval được tính bằng tpour.Encoder từ checkpoint trước khi tạo viewer; trình duyệt chỉ hiển thị và đổi query đã tính, không tự chạy huấn luyện. Ảnh là screenshot thật của viewer, không phải hình giao diện dựng bằng công cụ sinh ảnh.")
    r.page("10.1. Ảnh chụp demo: xếp hạng tài liệu")
    r.image("demo_02_retrieval.png", "Hình 5. So sánh top-5 của encoder toy trước train và CE + TRPO; màu xanh chỉ positive theo qrels.", width=6.1)
    r.p("Truy vấn demo hỏi về một tổ chức test. Cùng một corpus chứa cả bản cũ 2018 và bản mới 2021. Cột score là cosine vì encoder toy dùng L2 normalization. Tài liệu đúng phiên bản được xác định bằng qrels độc lập với cách sắp xếp. Hạng tài liệu đúng được tính trên toàn bộ 16 tài liệu, dù ảnh chỉ hiển thị top-5.")
    r.p("Ảnh làm rõ sự thay đổi của một ví dụ dễ quan sát, nhưng không chứng minh TRPO tốt hơn CE-only. Khi thuyết trình, quay lại bảng đối chứng toàn test để giữ diễn giải cân bằng. Tập giả lập có thay đổi quy trình rất rõ, chưa kiểm tra các biến động lịch sử khó hơn.")
    r.page("10.2. Ảnh chụp demo: log và kiểm tra")
    r.image("demo_03_checks.png", "Hình 6. Log kiểm tra và log demo được hiển thị từ đầu ra lệnh đã chạy; timestamp và provenance có trong evidence.", width=6.1)
    r.p("Ảnh kiểm tra giúp chứng minh kết quả không chỉ được điền tay vào slide. Có thể mở file tests.txt để thấy tên năm kiểm tra, kết quả và exit code. Log demo seed 42 ghi loss ở các bước thực tế. Config và lịch sử đầy đủ 150 bước được lưu riêng để kiểm tra sâu hơn.")
    r.p("Các ảnh không chứa terminal của máy người dùng; chúng là ảnh trình duyệt của trang bằng chứng. Điều này cho phép trình bày log rõ ràng mà vẫn giữ nguồn dữ liệu gốc trong repository. Trạng thái Contriever ở ảnh được ghi rõ là smoke test lịch sử, không phải benchmark chạy mới.")
    r.page("11. Nội suy tham số theo thời gian")
    r.h("11.1. Time vector trong bài báo")
    r.code("tau_t = theta_t - theta_base\ntheta_mid = theta_base + (1-alpha)*tau_left + alpha*tau_right\n          = (1-alpha)*theta_left + alpha*theta_right")
    r.p("Time vector là độ lệch tham số so với cùng mô hình nền. Đây không phải đơn thuần cộng một vector năm vào embedding của query. Khi hai mô hình được fine-tune từ cùng nền và cùng kiến trúc/tokenizer, nội suy trọng số có thể tạo encoder ở mốc trung gian [1]. α = 0 chọn mô hình trái; α = 1 chọn mô hình phải. Hệ số theo lịch, ví dụ khoảng 1/3 giữa 2018 và 2021 cho 2019, là một heuristic, không bảo đảm tối ưu trên mọi dữ liệu.")
    r.h("11.2. Điều kiện và kiểm tra trong code")
    r.p("Lệnh interpolate kiểm tra metadata model, max length, normalization, key/shape/dtype tensor, α trong [0,1] và sự bằng nhau của buffer không phải số thực. Nó không tự chứng minh hai checkpoint cùng mô hình nền. Người thực nghiệm phải giữ provenance, cùng seed khởi tạo hoặc cùng pretrained checkpoint và cùng tokenizer.")
    r.h("11.3. Điều đã và chưa kiểm chứng")
    r.p("Trong lượt triển khai trước đã tạo checkpoint nội suy giữa hai nhánh toy để kiểm tra thao tác tensor và ghi file. Hai nhánh đó cùng nền nhưng không đại diện hai mốc thời gian riêng biệt; vì thế không dùng thao tác này để khẳng định tổng quát hóa 2019/2020. Một thí nghiệm đúng cần huấn luyện riêng theo 2018 và 2021, đánh giá intermediate years bằng qrels phù hợp và so với checkpoint hai đầu.")
    r.h("11.4. Chi phí inference")
    r.p("Sau nội suy phải mã hóa lại corpus bằng encoder mới. Để đánh giá nhiều α, có thể cần nhiều lần tạo embeddings/index và nhiều checkpoint. Đó là chi phí thực tế ngoài phép nội suy tham số vốn đơn giản. Phần mixture-of-TPOUR cho dự đoán timestamp trong bài chưa được triển khai trong bộ thực hành.")
    r.page("12. Giới hạn và khác biệt với bài gốc")
    r.table(["Thành phần", "Bản thực hành", "Hệ quả khi diễn giải"], [["Reference", "Eval mode, momentum update sau optimizer", "Không đồng nhất toàn bộ dropout/thứ tự của upstream"], ["CE negatives", "Queue embeddings cũ; không thêm in-batch", "Phân phối logits khác code tác giả"], ["Temperature", "Dùng trong CE và TRPO", "Upstream xử lý hai nhánh khác nhau"], ["Queue khởi tạo", "Dữ liệu train thay vector ngẫu nhiên", "Bước đầu khác upstream"], ["Scheduler", "Learning rate cố định", "Không tái lập lịch linear + warmup"], ["Data mining", "Chia nửa + cosine nearest other-time", "Heuristic nhỏ, không pipeline Wikipedia đầy đủ"], ["Checkpoint", "Chỉ encoder và metadata", "Không resume optimizer/reference/queue"], ["Timestamp head", "Chưa triển khai", "Không có kết quả mixture-of-TPOUR"]])
    r.h("12.1. Những rủi ro khoa học của demo")
    r.p("Tín hiệu thời gian và tín hiệu từ vựng bị gắn chặt: digital/online luôn là năm mới. Một mô hình có thể học ‘ưu tiên digital’ thay vì hiểu quan hệ thời gian. Test khác thực thể nhưng cùng template không giải quyết vấn đề này. Ngoài ra, corpus test nhỏ và một positive/query làm chỉ số nhanh đạt trần.")
    r.h("12.2. Đánh giá lại nhận định ‘rất dễ kiểm chứng’")
    r.p("Dễ viết skeleton không đồng nghĩa dễ tái lập nghiên cứu. Chi phí đáng kể nằm ở snapshot thời gian, chọn các nội dung thay đổi có ý nghĩa, mining cùng chủ đề, chống rò rỉ và tạo temporal qrels. Repository cung cấp bước khởi đầu chạy được; không có bằng chứng so sánh đủ 11 bài để khẳng định TPOUR dễ nhất. Báo cáo không sử dụng nhận định xếp hạng đó như một kết luận đã kiểm chứng.")
    r.page("13. Hướng dẫn chạy lại và trình diễn")
    r.h("13.1. Lấy mã và cài môi trường")
    r.code("git clone `\n  https://github.com/lequangduccode/tpour-trpo-experiments.git\ncd tpour-trpo-experiments\npython -m venv .venv\n.venv\\Scripts\\python -m pip install -r requirements.txt")
    r.p("Ví dụ lệnh dùng PowerShell; dấu backtick ở cuối dòng nối lệnh sang dòng sau. Trên Linux/macOS dùng một dòng và thay đường dẫn Python môi trường ảo bằng .venv/bin/python. Mã tác giả nằm trực tiếp trong upstream/TPOUR; commit và hash gốc được ghi trong upstream/UPSTREAM.json, không cần lệnh submodule.")
    r.h("13.2. Chạy demo và kiểm tra")
    r.code(".venv\\Scripts\\python -m unittest test_tpour -v\n.venv\\Scripts\\python tpour.py demo --output outputs/demo\n.venv\\Scripts\\python tpour.py evaluate `\n  --data outputs/demo/data/test `\n  --checkpoint outputs/demo/tpour/encoder.pt `\n  --output outputs/demo/recheck.json")
    r.h("13.3. Trình diễn trong 2 phút")
    r.bullets(["Mở reports/demo_viewer.html: giới thiệu dữ liệu giả lập, quy mô và bảng đối chứng.", "Chuyển tab Truy xuất demo: chọn query, chỉ hạng và tài liệu đúng mốc; giải thích score.", "Chuyển tab Log & kiểm tra: chỉ kiểm tra đã đạt và các loss đã chạy.", "Kết thúc bằng giới hạn: chưa BEIR, chưa chứng minh lợi thế TRPO, cần dữ liệu thời gian khó hơn."])
    r.h("13.4. Tạo lại tài liệu")
    r.code(".venv\\Scripts\\python -m pip install -r requirements-report.txt\n.venv\\Scripts\\python tools/prepare_report_evidence.py\n.venv\\Scripts\\python tools/build_report.py\n.venv\\Scripts\\python tools/export_office.py")
    r.p("Chụp ảnh tự động dùng Playwright với Chrome cài sẵn. Công cụ export Office là bước riêng trên Windows có Word/PowerPoint. Nội dung Word và PowerPoint có thể chỉnh sửa trực tiếp; thông tin người thực hiện trên bìa được để trống để điền trước khi nộp.")
    r.page("14. Kế hoạch đánh giá mở rộng")
    r.h("14.1. Thiết kế temporal retrieval khó hơn")
    r.p("Thu thập các phiên bản cùng thực thể qua thời gian, đa dạng hóa loại cập nhật và template query. Negative cần cùng chủ đề nhưng sai phiên bản. Nên có truy vấn tường minh và ngầm định, cả những trường hợp không đơn thuần ưu tiên năm mới nhất. Tách theo thực thể/tài liệu và giữ mốc trung gian chưa dùng train để kiểm tra generalization.")
    r.table(["Thí nghiệm đề xuất", "Đối chứng / tiêu chí"], [["Baseline tiền huấn luyện", "Contriever gốc trên cùng corpus và qrels"], ["CE-only / CE+TRPO", "Cùng số bước, dữ liệu và ít nhất ba seed"], ["Ablation λ, β, queue", "Tách ảnh hưởng preference loss khỏi negative budget"], ["Mốc trung gian", "Hai checkpoint đầu/cuối và nhiều α, mã hóa corpus lại"], ["BEIR", "Full corpus, full judged queries, cùng dot-product protocol"], ["Chi phí", "RAM/VRAM, thời gian train, encode/index, inference/query"]])
    r.h("14.2. BEIR và temporal QA")
    r.p("BEIR kiểm tra retrieval tổng quát; script evaluate_beir.py dùng format gốc và metric thư viện [4]. Nó giữ đầy đủ corpus, tùy chọn query-limit chỉ phục vụ smoke test và có cờ partial_queries. Để chứng minh ưu tiên thời gian cần thêm SituatedQA/RealTimeQA hoặc tập temporal retrieval có tài liệu phiên bản và qrels đúng mốc. Không thay full corpus bằng tập con dễ hơn rồi so trực tiếp với kết quả công bố.")
    r.h("14.3. Kết luận")
    r.p("Bộ thực hành đã triển khai được các thành phần cốt lõi và tạo bằng chứng chạy thực tế trên CPU. Các kiểm tra và smoke test hỗ trợ kết luận về khả năng vận hành của pipeline. Demo cho thấy quy luật giả lập dễ học, nhưng đối chứng CE-only đạt trần như TPOUR nên chưa đo được lợi ích riêng của TRPO. Đóng góp phù hợp của báo cáo là phân tích phương pháp, xây dựng mã và xác định lộ trình kiểm chứng đầy đủ hơn.", True)
    r.page("Tài liệu tham khảo và nguồn bằng chứng")
    refs = [
        "[1] Kim, H.; Shim, J.; Kim, Y. J.; Bak, J. Temporal Preference Optimization for Unsupervised Retrieval. arXiv:2606.17664 (2026). " + PAPER + ". Cơ sở: mục 3 và phụ lục tái lập.",
        "[2] agwaBom/TPOUR. Repository triển khai chính thức. https://github.com/agwaBom/TPOUR. Mã được lưu trực tiếp trong upstream/TPOUR, tại commit gốc: " + meta["upstream_commit"] + ". Tệp src/moco.py và finetune.sh; provenance ở upstream/UPSTREAM.json.",
        "[3] Izacard và cộng sự. Towards Unsupervised Dense Information Retrieval with Contrastive Learning. https://arxiv.org/abs/2112.09118. Model card và mean pooling: https://huggingface.co/facebook/contriever.",
        "[4] Thakur và cộng sự. BEIR: A Heterogeneous Benchmark for Zero-shot Evaluation of Information Retrieval Models. https://arxiv.org/abs/2104.08663. Code: https://github.com/beir-cellar/beir.",
        "[5] Rafailov và cộng sự. Direct Preference Optimization: Your Language Model is Secretly a Reward Model. https://arxiv.org/abs/2305.18290.",
        "[6] Bộ thực hành của báo cáo. " + REPO + ". Tệp tpour.py, test_tpour.py, evaluate_beir.py; reports/evidence và results/contriever_smoke."
    ]
    for ref in refs:
        r.p(ref)
    r.h("Truy xuất nguồn kết quả")
    r.table(["Nội dung", "File bằng chứng"], [["Ba seed và chỉ số", "reports/evidence/summary.json, seed_42/43/44.json"], ["Loss", "reports/evidence/ce_only_train_log.json, tpour_train_log.json"], ["Log lệnh", "reports/evidence/tests.txt, demo_seed_42/43/44.txt"], ["Điểm từng tài liệu", "reports/evidence/retrieval_rankings.json"], ["Môi trường/hash/timestamp", "reports/evidence/provenance.json"], ["Contriever lịch sử", "results/contriever_smoke/*.json"], ["Ảnh chụp Chrome", "reports/assets/demo_01/02/03_*.png"]])
    r.p("Báo cáo này trích dẫn và diễn giải bài báo; các biểu đồ, sơ đồ và ảnh viewer do bộ thực hành tạo. Không chép nguyên bảng benchmark của tác giả vào bảng kết quả tự đo. Ngày tạo bằng chứng mới được lấy theo múi giờ Asia/Bangkok.")
    r.page("Phụ lục: câu hỏi khi thuyết trình")
    questions = [
        ("TRPO khác ranking loss thông thường ở đâu?", "TRPO so sánh chênh lệch preferred–rejected của encoder hiện tại với chênh lệch tương ứng của reference; reference được cập nhật momentum. Bỏ phần reference sẽ thay đổi objective."),
        ("Dữ liệu không giám sát có thật sự không cần thời gian?", "Không cần query–document labels thủ công cho train theo cách đề xuất, nhưng vẫn cần corpus có mốc thu thập và khác biệt nội dung có ý nghĩa. Demo tự dựng tín hiệu nên chỉ là minh họa."),
        ("Vì sao đạt nDCG 1,0 mà chưa chứng minh ưu thế?", "Vì CE-only cũng 1,0, test rất nhỏ và cùng template. Đây là hiệu ứng trần; phải có negative khó hơn và benchmark thật."),
        ("Demo có dùng Contriever không?", "Bảng chính dùng toy encoder; Contriever thật chỉ có smoke test riêng 2 bước. Không trộn hai loại kết quả."),
        ("Có thể kết luận mô hình hiểu thời gian không?", "Chưa. Nó có thể học từ vựng digital/online. Cần truy vấn phản ví dụ, cặp cùng chủ đề và mốc trung gian chưa học."),
        ("Nội suy có phải nội suy embeddings?", "Không. Bài dùng nội suy tham số giữa các encoder fine-tune từ cùng nền, sau đó mã hóa lại query/corpus."),
        ("Tại sao cần reference và queue?", "Reference tạo mốc so ưu tiên và embeddings thay đổi chậm; queue cung cấp negative cũ để mở rộng contrastive set. Hai vai trò liên quan nhưng khác nhau."),
        ("Bước tiếp theo quan trọng nhất?", "Giữ cùng protocol cho Contriever gốc, CE-only và TPOUR trên dữ liệu temporal thực đã chống rò rỉ; báo cáo cả chỉ số retrieval và chi phí tính toán.")]
    for q, a in questions:
        r.h(q)
        r.p(a)
    doc.save(OUT / "Bao_cao_chi_tiet_TPOUR_TRPO.docx")
    return refs


BG = "F3F6FB"
INK = "14243D"
BLUE = "3767E8"
GREEN = "159987"
MUTED = "63748A"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = PI(13.333), PI(7.5)
        self.notes = []

    def text(self, slide, text, x, y, w, h, size=22, color=INK, bold=False):
        box = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name, p.font.size = "Arial", PP(size)
            p.font.color.rgb = PC.from_string(color)
            p.font.bold = bold
            p.space_after = PP(10)
        return box

    def shape(self, slide, x, y, w, h, fill, rounded=True):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                        PI(x), PI(y), PI(w), PI(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PC.from_string(fill)
        shape.line.fill.background()
        return shape

    def slide(self, title, subtitle, notes):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PC.from_string(BG)
        self.shape(slide, .45, .48, .1, .45, BLUE, False)
        self.text(slide, "TPOUR / TRPO • NGHIÊN CỨU & THỰC HÀNH", .72, .27, 11.5, .32, 11, MUTED, True)
        self.text(slide, title, .72, .83, 12, .72, 31, INK, True)
        self.text(slide, subtitle, .72, 1.57, 12, .52, 16, MUTED)
        n = len(self.prs.slides)
        self.text(slide, "12/10/2026   |   Demo CPU • Chưa tái lập đầy đủ bài báo", .72, 7.06, 11.5, .22, 10, MUTED)
        self.text(slide, f"{n:02d}", 12.15, 7.01, .5, .3, 13, BLUE, True)
        slide.notes_slide.notes_text_frame.text = notes
        self.notes.append((title, notes))
        return slide

    def bullets(self, slide, items, x=.9, y=2.4, w=11.5, size=24, gap=.9):
        for i, item in enumerate(items):
            self.shape(slide, x, y+i*gap+.12, .1, .1, BLUE, False)
            self.text(slide, item, x+.3, y+i*gap, w-.3, gap-.12, size)

    def card(self, slide, title, body, x, y, w, h, color=BLUE):
        self.shape(slide, x, y, w, h, "FFFFFF")
        self.text(slide, title, x+.25, y+.22, w-.5, .55, 21, color, True)
        self.text(slide, body, x+.25, y+.92, w-.5, h-1.0, 20)

    def image(self, slide, name, x, y, w, h):
        from PIL import Image
        path = ASSETS / name
        with Image.open(path) as img:
            ratio = img.width/img.height
        width = min(w, h*ratio)
        height = width/ratio
        slide.shapes.add_picture(str(path), PI(x+(w-width)/2), PI(y+(h-height)/2), width=PI(width), height=PI(height))


def build_slides(summary, ranking):
    d = Deck()
    s = d.slide("TPOUR / TRPO", "Tối ưu ưu tiên thời gian cho hệ thống truy xuất không giám sát",
        "Thời lượng gợi ý: 12–15 phút, cộng 2 phút demo. Giới thiệu rằng đây là báo cáo phân tích phương pháp và bản thực hành PyTorch. Trình bày hai lớp bằng chứng: kết quả giả lập và smoke test Contriever thật. Ngay đầu nói rõ chưa tái lập đầy đủ benchmark của bài báo. Điền tên, lớp và mã sinh viên trước khi thuyết trình.")
    d.text(s, "Từ thuật toán → mã chạy được → bằng chứng kiểm tra", .9, 2.45, 11.5, .8, 29, BLUE, True)
    d.card(s, "Phân tích", "TRPO • MoCo\nMomentum reference", .9, 3.6, 3.7, 2)
    d.card(s, "Thực hành", "PyTorch trên CPU\nToy + Contriever", 4.8, 3.6, 3.7, 2)
    d.card(s, "Kiểm chứng", "Đối chứng chỉ CE\nẢnh demo & log thật", 8.7, 3.6, 3.7, 2)
    d.text(s, "Người trình bày: ........................................   |   Lớp: ................", .9, 6.35, 11.5, .35, 17, MUTED)
    s = d.slide("01 / Đúng chủ đề vẫn có thể sai thời điểm", "Temporal alignment là một phần của độ liên quan [1]",
        "Dành khoảng 45 giây. Dùng ví dụ tổ chức đã chuyển từ phục vụ ở quầy sang cổng trực tuyến. Cả hai phiên bản đều cùng chủ đề và cùng tên tổ chức, nên dense retrieval có thể thấy chúng gần nhau. Truy vấn hiện tại cần bản phù hợp thời điểm. Đây là ví dụ giả định của demo; không phải dữ kiện về một tổ chức có thật.")
    d.card(s, "Truy vấn", "How does Organization32\ncurrently provide its service?", .9, 2.35, 11.5, 1.9)
    d.card(s, "2018 / Lệch mốc", "Paper forms\nOffice desk service", .9, 4.45, 5.5, 1.95, MUTED)
    d.card(s, "2021 / Đúng mốc đích", "Digital portal\nOnline mobile service", 6.7, 4.45, 5.7, 1.95, GREEN)
    s = d.slide("02 / TPOUR thêm tín hiệu ưu tiên vào retrieval", "Cốt lõi: học ngữ nghĩa + ưu tiên corpus đúng mốc [1, 2]",
        "Dành 45 giây. Giải thích TPOUR là framework, TRPO là objective preference trong framework đó. Query có thể nêu năm hoặc chỉ nói hiện tại. Mô hình được căn chỉnh theo mốc huấn luyện, không tự biết đồng hồ máy. Corpus-level temporal signal không đồng nghĩa hoàn toàn không cần dữ liệu về thời gian. Bản thực hành không triển khai mô hình phân loại timestamp mixture-of-TPOUR.")
    d.bullets(s, ["Contrastive loss: giữ độ liên quan ngữ nghĩa", "TRPO: ưu tiên đúng mốc hơn lệch mốc", "Reference cập nhật momentum + queue negative", "Không cần reward model riêng; vẫn cần cặp thời gian tốt"], size=25)
    s = d.slide("03 / Hai encoder, một hàng đợi", "Luồng bản thực hành: optimizer → momentum update → enqueue",
        "Dành khoảng một phút. Theo các mũi tên từ bộ ba tới hai encoder. Encoder hiện tại nhận gradient; reference eval/no_grad và cập nhật bằng momentum. Query/current positive được so với queue cũ trong CE; TRPO dùng chênh điểm hai encoder. Sau optimizer, reference cập nhật và keys được ghi vào queue. Sơ đồ tự vẽ theo code và phương pháp, không phải screenshot giao diện hoặc figure sao chép.")
    d.image(s, "architecture.png", .8, 2.2, 11.8, 4.35)
    s = d.slide("04 / TRPO so sánh chênh lệch với reference", "TRPO = Temporal Retrieval Preference Optimization [1]",
        "Dành một phút. Delta_theta là điểm preferred trừ rejected của current encoder, Delta_ref là hiệu tương ứng của reference. Lấy chênh lệch hai delta, nhân beta và log-sigmoid. Khi current tăng ưu tiên so với reference, loss giảm. Không nhầm TRPO của bài này với Trust Region Policy Optimization. Tại hai gap bằng nhau, loss ln2 nhưng gradient không bằng không. Reference không nhận gradient trực tiếp.")
    d.card(s, "Chênh lệch ưu tiên", "Δθ = Sθ(Q,Dđúng) − Sθ(Q,Dlệch)\nΔref = Sref(Q,Dđúng) − Sref(Q,Dlệch)", .9, 2.3, 11.5, 1.8)
    d.shape(s, .9, 4.4, 11.5, 1.15, INK)
    d.text(s, "L_TRPO = − mean[log σ(β · (Δθ − Δref))]", 1.2, 4.73, 10.9, .55, 28, "FFFFFF", True)
    d.text(s, "Dùng log-sigmoid ổn định số • Reference detach • β = 0,3", 1, 5.95, 11, .5, 20, MUTED)
    s = d.slide("05 / Loss tổng và cơ chế MoCo", "Tham số λ cân bằng hai objective; không là tỷ lệ mẫu",
        "Dành 45 giây. Đọc công thức loss tổng, chỉ CE khi lambda bằng1. Với lambda0.925, trọng số TRPO là0.075 nhưng ảnh hưởng còn phụ thuộc thang loss và gradient. Temperature mặc định0.05. Momentum lớn khiến reference thay đổi chậm. Queue chứa negative cũ nên tiết kiệm yêu cầu batch đồng thời; queue vẫn có chi phí bộ nhớ và phép nhân. Bản gọn khác code tác giả ở in-batch negatives và temperature CE.")
    d.card(s, "Loss tổng", "L = λ · LCE + (1 − λ) · LTRPO\nλ = 0,925   |   T = 0,05", .9, 2.45, 11.5, 1.8)
    d.card(s, "Reference", "θref ← m · θref + (1 − m) · θ\nm = 0,9999", .9, 4.5, 5.5, 1.8, GREEN)
    d.card(s, "Queue", "Lưu embeddings tài liệu reference\nCho negative ở các batch sau", 6.7, 4.5, 5.7, 1.8)
    s = d.slide("06 / Dữ liệu nhỏ để kiểm tra luồng chạy", "Tách thực thể train/test; template vẫn giống nhau",
        "Dành 45 giây. Nêu 40 tổ chức giả định, train tổ chức0–31, test32–39. Mỗi query test có một relevant document: bản mới đúng tổ chức. Dữ liệu mốc và qrels được dựng trong demo; không gọi nó là tái lập unsupervised learning trên Wikipedia. Hash word embedding không phải Contriever. Template chung và digital/online gắn mốc khiến bài toán rất dễ và cần nói rõ trước khi trình bày điểm cao.")
    d.card(s, "Train", "32 bộ ba\nQ + D2021 + D2018", .9, 2.4, 3.65, 2)
    d.card(s, "Test", "8 query • 16 documents\n1 positive / query", 4.8, 2.4, 3.65, 2)
    d.card(s, "Toy encoder", "4096 × 64 embedding\nMean pool + L2", 8.7, 2.4, 3.7, 2)
    d.text(s, "Giới hạn dữ liệu: ‘digital/online’ luôn đi với năm mới.\nMô hình có thể học từ vựng thay vì hiểu thời gian.", 1, 4.95, 11.3, 1.2, 26, INK, True)
    s = d.slide("07 / Thiết kế đối chứng", "Cùng nền khởi tạo, dữ liệu, seed và số bước",
        "Dành 40 giây. Ba nhánh là toy chưa train, chỉ CE, CE+TRPO. Hai nhánh train được sao chép từ encoder ban đầu rồi dùng cùng cấu hình tương ứng theo từng seed. Chỉ CE có lambda1. Ba seed42,43,44 là ba tiến trình thật chạy bổ sung. AdamW learning rate0.01 thuộc toy demo, không phải thiết lập fine-tune Contriever. Các chỉ số trên test8query không đủ suy luận thống kê mạnh.")
    d.card(s, "Chỉ CE", "λ = 1,0\nTRPO ghi log, không tối ưu", .9, 2.4, 5.5, 1.9, GREEN)
    d.card(s, "CE + TRPO", "λ = 0,925\nObjective ưu tiên thời gian", 6.7, 2.4, 5.7, 1.9)
    d.bullets(s, ["150 bước • batch 8 • queue 256 • learning rate 0,01", "Seed 42, 43, 44 • CPU • không xóa từ ở demo"], y=4.85, size=24)
    s = d.slide("08 / Kết quả: cả hai nhánh đã train đạt trần", "Seed 42 • toàn bộ 8 query test • thang điểm 0–1",
        "Dành một phút. Biểu đồ lấy từ JSON đầu ra thực tế. Toy chưa train nDCG@5 khoảng0.3028, Recall@5 0.625 và MRR@5 0.2. Chỉ CE và CE+TRPO đều1.0. Vì vậy chỉ có thể nói pipeline học được bài toán giả lập, không thể nói TRPO vượt CE. Đừng nhầm untrained với pretrained Contriever. Demo tạo qrels một positive/query nên Recall ở đây giống hit-rate; trong dataset nhiều positive khái niệm khác nhau.")
    d.image(s, "metrics.png", .9, 2.2, 11.5, 3.9)
    d.text(s, "Δ(CE + TRPO − chỉ CE) = 0  →  chưa thấy lợi ích riêng của TRPO", 1, 6.35, 11.3, .5, 21, BLUE, True)
    s = d.slide("09 / Lặp với 3 seed để kiểm tra độ nhạy", "Dữ liệu vẫn giả lập; các seed không thay thế benchmark thực",
        "Dành 40 giây. Đọc bảng các lần chạy thật, so untrained với hai nhánh train và deltaTRPO–CE. Nêu rằng khi mọi nhánh train đạt trần trên test nhỏ, nhiều seed vẫn không giúp chứng minh khác biệt. Muốn đo lợi ích phải tăng độ khó và đa dạng corpus/query, chống rò rỉ và dùng test độc lập có temporal qrels. Logs từng seed và hash code có trong reports/evidence.")
    for i, seed in enumerate(("42", "43", "44")):
        m = summary[seed]
        d.card(s, "Seed " + seed,
               f"Chưa train: {m['untrained']['nDCG@5']:.3f}\nChỉ CE: {m['ce_only']['nDCG@5']:.3f}\nCE+TRPO: {m['tpour']['nDCG@5']:.3f}", .9+i*3.9, 2.5, 3.65, 2.7)
    d.text(s, "nDCG@5 • cùng corpus / query / protocol đánh giá", 1, 5.65, 11.3, .7, 25, MUTED)
    s = d.slide("10 / Quan sát loss và tính đúng", "Log seed 42 + 5 kiểm tra tự động thành công",
        "Dành 50 giây. Đồ thị trục log lấy loss từng bước. CE-only vẫn ghi TRPO loss nhưng không tối ưu nó trực tiếp. Không so loss tổng thấp hơn của hai objective khác trọng số như chất lượng tốt hơn. Nêu năm kiểm tra: loss/gradient, reference detach, queue/momentum, metric tính tay, input tác giả. Tests xác nhận những bất biến cụ thể, không chứng minh mọi tình huống hoặc generalization.")
    d.image(s, "loss.png", .9, 2.2, 11.5, 3.8)
    d.text(s, "5/5 đạt: gradient • reference • queue/momentum • metrics • input", 1, 6.35, 11.3, .5, 20, GREEN, True)
    s = d.slide("11 / Contriever thật đã chạy trên CPU", "Smoke test ngày 10/10/2026 • không phải kết quả benchmark",
        "Dành 45 giây. Nói rõ đây là log lịch sử đã chạy trong lượt triển khai trước. Model facebook/contriever với 8 mẫu đầu từ dataset tác giả,2bước,batch2,maxlength64,queue8,lr1e-6. Checkpoint438MB đã lưu và nạp lại. Loss bước2 cao hơn bước1 do batch và queue thay đổi, hai điểm không cho thấy hội tụ. Nạp lại trên synthetic corpus không chứng minh hai bước fine-tune cải thiện vì chưa có baseline pretrained trên cùng tập.")
    d.card(s, "Cấu hình", "8 mẫu • 2 bước • batch 2\nMax length 64 • queue 8\nLearning rate 1e-6", .9, 2.5, 5.5, 2.7)
    d.card(s, "Bằng chứng", "Forward / backward thành công\nLưu & nạp checkpoint ~438 MB\nLoss: 0,04648 → 0,27002", 6.7, 2.5, 5.7, 2.7, GREEN)
    d.text(s, "Kết luận: đường chạy mô hình thật hoạt động. Chưa đánh giá retrieval đầy đủ.", 1, 5.75, 11.3, .8, 24, BLUE, True)
    s = d.slide("12 / Ảnh demo: từ query đến xếp hạng", "Ảnh chụp Chrome từ viewer; điểm tính bằng checkpoint toy",
        "Dành khoảng một phút, hoặc chuyển sang mở viewer trực tiếp. Chỉ query test, corpus2018/2021, hạng đúng và cosine score. Hàng xanh theo qrels, không tô dựa trên model prediction. Ví dụ được chọn là baseline có hạng đúng kém nhất để dễ quan sát, nên mang tính minh họa. Nhắc rằng bảng đối chứng toàn test vẫn không có lợi ích riêngTRPO. Viewer là kết quả tính trước, không huấn luyện trong browser.")
    d.image(s, "demo_02_retrieval_detail.png", .75, 2.35, 11.85, 4.1)
    s = d.slide("13 / Ảnh demo: log có thể kiểm tra", "Bằng chứng lệnh chạy, exit code, seed và provenance",
        "Dành 40 giây. Ảnh Chrome này hiển thị log lệnh đã chạy lại. Nếu thầy yêu cầu xem bằng chứng, mở reports/evidence/tests.txt, demo_seed_42.txt, summary.json và provenance.json. Các ảnh và biểu đồ được tạo từ file kết quả trong repository. Phân biệt kiểm tra mới ngày 12/10 với log Contriever ngày 10/10. Có thể chạy lại các script tạo bằng chứng và tài liệu.")
    d.image(s, "demo_03_checks_detail.png", .75, 2.35, 11.85, 4.1)
    s = d.slide("14 / Nội suy thời gian là nội suy trọng số", "Cùng mô hình nền, kiến trúc và tokenizer [1]",
        "Dành 45 giây. Time vector là theta_t trừ theta_base; phép cộng triệt tiêu base để thành nội suy hai checkpoint. Đây là nội suy trọng số, không phải cộng year embedding. Điều kiện cần là cùng pretrained checkpoint nền. Sau nội suy phải encode corpus lại. Bản gọn chỉ kiểm tra thao tác tensor; nội suy CE-only và TPOUR toy không tương ứng hai năm 2018 và 2021, nên chưa chứng minh tổng quát hóa thời gian.")
    d.shape(s, .9, 2.5, 11.5, 1.45, INK)
    d.text(s, "θmid = (1 − α) · θleft + α · θright", 1.2, 2.98, 10.9, .6, 31, "FFFFFF", True)
    d.bullets(s, ["Time vector = độ lệch tham số so với cùng nền", "Mã hóa corpus lại bằng encoder nội suy", "Chưa có kiểm chứng mốc trung gian trong demo hiện tại"], y=4.45, gap=.75, size=23)
    s = d.slide("15 / Giới hạn là một phần của kết quả", "Skeleton chạy được chưa đồng nghĩa tái lập bài báo",
        "Dành một phút. Liệt kê template chung, tín hiệu từ vựng dễ học, corpus nhỏ, chưa BEIR và khác biệt CE với upstream. Phần khác biệt được mô tả trong Word, gồm reference eval, negative queue, temperature, scheduler và queue warmup. Chưa có resume toàn bộ checkpoint, timestamp head hoặc thí nghiệm mốc trung gian. Không khẳng định dễ nhất trong 11 bài vì chưa có phép so sánh đầy đủ. Khi đối chứng không có chênh lệch, cần báo cáo đúng kết quả đó.")
    d.bullets(s, ["Toy encoder + template chung → bài toán dễ đạt trần", "Chỉ CE cũng đạt tối đa → chưa chứng minh TRPO tốt hơn", "2 bước Contriever → chỉ xác nhận khả năng chạy", "Chưa BEIR / temporal QA; bản gọn khác upstream"], size=24)
    s = d.slide("16 / Bước tiếp theo: đo giá trị gia tăng thật", "Giữ cùng protocol cho Contriever gốc, CE-only và TPOUR",
        "Dành 40 giây. Ưu tiên dataset cùng chủ đề nhưng khác phiên bản với nội dung thực và temporal qrels. Tách theo document/entity, loại near-duplicate. So pretrained Contriever, CE-only và TPOUR cùng ngân sách; dùng ít nhất ba seed, thêm khoảng bất định khi test đủ lớn. BEIR đo retrieval tổng quát; temporal QA đo căn chỉnh thời gian. Đánh giá nội suy riêng với checkpoint hai đầu và mốc trung gian. Đo bộ nhớ, train và inference để nhận định chi phí triển khai.")
    d.card(s, "Dữ liệu", "Snapshot thực theo mốc\nNegative cùng chủ đề\nChống rò rỉ / near-duplicate", .9, 2.45, 3.65, 2.65)
    d.card(s, "Đánh giá", "Baseline Contriever gốc\nCE-only / CE+TRPO\nBEIR + temporal qrels", 4.8, 2.45, 3.65, 2.65)
    d.card(s, "Chi phí", "RAM / VRAM\nTrain, encode, indexing\nInference mỗi query", 8.7, 2.45, 3.7, 2.65)
    d.text(s, "Thí nghiệm nội suy: hai mốc train → một mốc test chưa học", 1, 5.75, 11.3, .65, 26, BLUE, True)
    s = d.slide("17 / Kết luận và hỏi đáp", "Đã có mã, tài liệu, ảnh demo và bằng chứng có thể chạy lại",
        "Dành 40 giây rồi hỏi đáp. Kết luận rõ: đã triển khai và chạy TRPO, reference và queue; có đối chứng và smoke test Contriever. Chưa chứng minh ưu thế retrieval riêng của TRPO. Nếu hỏi vì sao nDCG bằng 1 mà chưa tốt hơn, trả lời CE-only cũng bằng 1 và dataset đã đạt trần. Nếu hỏi hiện tại là năm nào, giải thích mốc căn chỉnh của mô hình khác đồng hồ máy. Nguồn bài báo, code tác giả và output/log đều được dẫn ở repository public.")
    d.bullets(s, ["Pipeline PyTorch chạy trên CPU; kiểm tra và checkpoint đạt", "Có đối chứng: điểm demo chưa cho thấy TRPO vượt CE", "Cần dữ liệu thời gian thực để kiểm chứng nghiên cứu"], size=25, gap=.92)
    d.text(s, "Nguồn: arxiv.org/abs/2606.17664 • github.com/agwaBom/TPOUR\nRepository: github.com/lequangduccode/tpour-trpo-experiments", 1, 5.7, 11.3, .85, 17, MUTED)
    s = d.slide("18 / Tài liệu tham khảo", "Nguồn thuật toán, mô hình, benchmark và bằng chứng tự đo",
        "Slide nguồn có thể để ở cuối phần hỏi đáp. Các số liệu demo đều truy xuất được trong reports/evidence. Bài báo chính là nguồn cho objective và time vector; code tác giả là nguồn đối chiếu triển khai. Contriever, BEIR và DPO có nguồn nghiên cứu gốc riêng. Báo cáo Word chứa thông tin trích dẫn đầy đủ hơn và commit upstream cụ thể.")
    d.bullets(s, ["[1] Kim và cộng sự, TPOUR (2026) — arxiv.org/abs/2606.17664", "[2] Code chính thức — github.com/agwaBom/TPOUR", "[3] Contriever — arxiv.org/abs/2112.09118", "[4] BEIR — arxiv.org/abs/2104.08663", "[5] DPO — arxiv.org/abs/2305.18290"], y=2.25, gap=.65, size=20)
    d.text(s, "Bằng chứng tự đo: reports/evidence • results/contriever_smoke\nRepo public: github.com/lequangduccode/tpour-trpo-experiments", 1, 5.85, 11.3, .8, 18, MUTED)
    d.prs.core_properties.title = "TPOUR/TRPO — Phân tích và thử nghiệm"
    d.prs.core_properties.author = ""
    d.prs.save(OUT / "Slide_thuyet_trinh_TPOUR_TRPO.pptx")
    notes_doc = Document()
    notes_doc.styles["Normal"].font.name = "Times New Roman"
    notes_doc.styles["Normal"].font.size = Pt(12)
    notes_doc.add_heading("Kịch bản thuyết trình TPOUR / TRPO", 0)
    notes_doc.add_paragraph("Thời lượng: 12–15 phút và 2 phút demo. Ghi chú cũng có trong Notes của từng slide PowerPoint. Điền tên/lớp trên slide đầu trước khi nộp.")
    for i, (title, note) in enumerate(d.notes, 1):
        notes_doc.add_heading(f"Slide {i:02d} — {title}", 1)
        notes_doc.add_paragraph(note)
    notes_doc.add_heading("Demo trực tiếp", 1)
    notes_doc.add_paragraph("Mở reports/demo_viewer.html bằng Chrome/Edge. Tổng quan → Truy xuất demo (đổi query) → Log & kiểm tra. Nếu không mở được viewer, dùng ba ảnh trong reports/assets. Không mô tả viewer như hệ thống search thời gian thực; nó trình bày các kết quả đã tính bằng checkpoint.")
    notes_doc.save(OUT / "Kich_ban_thuyet_trinh_TPOUR_TRPO.docx")
    return len(d.prs.slides)


def main():
    summary = j(EV / "summary.json")
    meta = j(EV / "provenance.json")
    ranking = j(EV / "retrieval_rankings.json")
    architecture()
    build_word(summary, meta, ranking)
    slides = build_slides(summary, ranking)
    print(f"Generated report, {slides} slides and speaker script in {OUT}")


if __name__ == "__main__":
    main()
