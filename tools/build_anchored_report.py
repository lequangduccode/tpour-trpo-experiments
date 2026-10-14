"""Write a Vietnamese Word supplement and six plain slides from verified GPU evidence."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI,Pt as PP
from pptx.dml.color import RGBColor


def read(path): return json.loads(Path(path).read_text(encoding='utf-8'))


def build(root,output):
    root=Path(root); output=Path(output); output.mkdir(parents=True,exist_ok=True)
    groups=read(root/'demo/summary.json')
    comparisons=[(p.name,read(p)) for p in sorted(root.glob('compare_tpour_seed*_anchored_seed*.json'))]
    if len(comparisons)!=3: raise ValueError('Need all three paired seeds before writing final conclusions')
    doc=Document(); section=doc.sections[0]
    section.page_width=Inches(8.27); section.page_height=Inches(11.69)
    for side in ['top_margin','bottom_margin','left_margin','right_margin']: setattr(section,side,Inches(1))
    for name in ['Normal','Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Times New Roman'; doc.styles[name].font.size=Pt(13 if name=='Normal' else 16)
    doc.add_heading('CẢI TIẾN TPOUR VÀ DEMO KIỂM CHỨNG',0)
    doc.add_paragraph('Phụ lục thực nghiệm của nhóm — fixed-teacher retention. Tài liệu này bổ sung cho báo cáo TPOUR trước đó, không thay thế nội dung bài báo gốc.')
    doc.add_heading('1. Vấn đề và thành phần mới',1)
    doc.add_paragraph('Thử nghiệm ban đầu dùng Wikipedia thật cho thấy CE+TRPO tốt hơn chỉ CE nhưng giảm nDCG@10 so với Contriever gốc. Nhóm đề xuất giữ thêm kiến thức của model khởi tạo bằng một teacher đóng băng. Đây là biến thể cải tiến bản thực hành, không tuyên bố phát minh kỹ thuật knowledge distillation.')
    doc.add_paragraph('TRPO gốc vẫn có current encoder, momentum reference và queue. Biến thể thêm teacher cố định: teacher không nhận gradient, không cập nhật momentum. Loss KL giữ phân bố xếp hạng trên cùng candidate list; relative embedding MSE giữ cả hướng và độ lớn embedding cho dot-product inference. Teacher chỉ dùng khi train, không trộn điểm baseline vào kết quả demo.')
    equation=doc.add_paragraph(); math=OxmlElement('m:oMath')
    def run(text):
        r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text; r.append(t); return r
    def sub(base,index):
        s=OxmlElement('m:sSub'); e=OxmlElement('m:e'); e.append(run(base)); b=OxmlElement('m:sub'); b.append(run(index)); s.extend([e,b]); return s
    for item in [sub('ℒ','new'),run(' = λ '),sub('ℒ','CE'),run(' + (1 − λ) '),sub('ℒ','TRPO'),run(' + γ ('),sub('ℒ','KL'),run(' + η '),sub('ℒ','rep'),run(')')]: math.append(item)
    equation._p.append(math)
    doc.add_paragraph('λ=0,925; γ=10; η=1. Retention temperature=0,1. KL(teacher || student) dùng cosine scores, tính bằng FP32. Representation loss bằng squared embedding error chia squared teacher norm, lấy trung bình trên query và tài liệu. Các giá trị cố định trước khi xem lần test mới.')
    doc.add_heading('2. Thiết kế kiểm chứng',1)
    doc.add_paragraph('Cùng dữ liệu: 2.543 mẫu train, 449 preference held-out từ mẫu Wikipedia tác giả. Tách theo văn bản nguyên dạng; chưa loại mọi trùng bài/đoạn trích. Đánh giá đủ 300 SciFact test queries và 5.183 tài liệu; không dùng test để lựa chọn gamma hoặc checkpoint.')
    doc.add_paragraph('Seed 42/43/44, mỗi nhánh 1.000 lần thử cập nhật; batch4, queue4096, lr1e-6, warmup100, max_length256, word deletion0,1, AMP và gradient checkpointing. Mỗi nhánh dùng một T4. Số cập nhật thực có thể khác khi AMP bỏ gradient không hữu hạn; bảng chi phí phải thể hiện điều này.')
    doc.add_paragraph('Đối chứng gồm Contriever gốc, TRPO gốc, anchored; ablation seed42 chỉ KL (η=0). Ba seed báo mean và sample SD; bootstrap ghép cặp theo query, 5.000 lần. CI theo query không thay thế bất định theo seed.')
    doc.add_heading('3. Kết quả thực',1)
    table=doc.add_table(rows=1,cols=4); table.style='Table Grid'
    for c,t in zip(table.rows[0].cells,['Nhánh','nDCG@10 mean','SD','Số lần train']): c.text=t
    for name in ['baseline','tpour','ranking_only','anchored']:
        s=groups[name]
        for c,t in zip(table.add_row().cells,[name,f'{s["mean"]:.4f}',f'{s["sample_sd"]:.4f}' if s['sample_sd'] is not None else '—',str(s['n'])]): c.text=t
    delta=groups['anchored']['mean']-groups['tpour']['mean']; baseline=groups['anchored']['mean']-groups['baseline']['mean']
    doc.add_paragraph(f'Chênh lệch mean anchored − TRPO: {delta:+.5f}; anchored − Contriever gốc: {baseline:+.5f}. Độ lớn và dấu của cả hai chênh lệch cần được trình bày, không chỉ chọn nhánh thuận lợi.')
    for name,c in comparisons: doc.add_paragraph(f'{name}: delta nDCG@10={c["mean_delta"]:+.5f}, CI95% [{c["ci95"][0]:+.5f}; {c["ci95"][1]:+.5f}]. Tốt hơn/kém hơn/không đổi: {c["positive_queries"]}/{c["negative_queries"]}/{c["tied_queries"]}.')
    doc.add_picture(str(root/'demo/01_full_test.png'),width=Inches(6.1))
    doc.add_heading('4. Chi phí và kiểm tra độ tin cậy',1)
    table=doc.add_table(rows=1,cols=4); table.style='Table Grid'
    for c,t in zip(table.rows[0].cells,['Run','Cập nhật thực','Train (giây)','Peak allocated GiB']): c.text=t
    for path in sorted(root.glob('*/completed.json')):
        c=read(path)
        for cell,t in zip(table.add_row().cells,[path.parent.name,str(c['optimizer_updates']),f'{c["total_seconds"]:.1f}',f'{c["peak_gpu_bytes"]/2**30:.2f}']): cell.text=t
    doc.add_paragraph('Thời gian trên chỉ đo hàm train, không gồm tải model, warmup queue và evaluate. Peak allocated không phải tổng VRAM thiết bị. 14 kiểm tra code/toán đã qua. Cần đối chiếu metrics từ rankings, qrels, hash dữ liệu, config và code commit; log thật nằm cùng kết quả.')
    doc.add_heading('5. Demo truy xuất và cách thuyết trình',1)
    doc.add_paragraph('Mở biểu đồ toàn bộ test trước, giải thích teacher cố định và objective, rồi trình bày query đầu tiên cố định. Hai ảnh tiếp theo là trường hợp cải thiện nhiều nhất và giảm nhiều nhất, được chọn từ kết quả để minh họa; không phải đánh giá khách quan thay cho toàn bộ test. Có thể chạy lại tools/demo_anchored_results.py từ run.json mà không train lại.')
    doc.add_picture(str(root/'demo/03_ranking_demo.png'),width=Inches(6.1))
    doc.add_heading('6. Giới hạn và nguồn',1)
    doc.add_paragraph('SciFact kiểm tra general retrieval, không chứng minh temporal alignment. Nhãn held-out là mined preference, chưa có temporal QA gold độc lập. Teacher có thể giữ thiên lệch/lỗi của model gốc. Hệ số thử nghiệm chưa tối ưu; ablation chỉ một seed. Không gọi kết quả là tái lập đầy đủ hoặc tính mới đã được peer review.')
    doc.add_paragraph('Li & Hoiem, Learning without Forgetting: https://arxiv.org/abs/1606.09282. Weighted KL-Divergence for Document Ranking Model Refinement: https://arxiv.org/abs/2406.05977. Bài TPOUR được chọn và source tác giả lưu ở upstream/TPOUR. Notebook: https://www.kaggle.com/code/vananh1911/tpour-retention-improvement-demo.')
    doc.save(output/'Bao_cao_cai_tien_TPOUR_retention.docx')
    pres=Presentation(); pres.slide_width=PI(13.333); pres.slide_height=PI(7.5)
    def slide(title,lines=None,picture=None):
        s=pres.slides.add_slide(pres.slide_layouts[6]); box=s.shapes.add_textbox(PI(.6),PI(.35),PI(12.1),PI(.75))
        p=box.text_frame.paragraphs[0]; p.text=title; p.font.size=PP(28); p.font.bold=True; p.font.color.rgb=RGBColor.from_string('943842')
        if lines:
            box=s.shapes.add_textbox(PI(.75),PI(1.5),PI(11.8),PI(5.2)); frame=box.text_frame; frame.word_wrap=True
            for i,line in enumerate(lines):
                p=frame.paragraphs[0] if i==0 else frame.add_paragraph(); p.text=line; p.font.size=PP(23); p.space_after=PP(22)
        if picture: s.shapes.add_picture(str(picture),PI(1.3),PI(1.35),width=PI(10.7))
        return s
    slide('Vì sao cần cải tiến?', ['Fine-tune Wikipedia nhỏ làm giảm SciFact so với Contriever gốc.','Giả thuyết: thêm teacher cố định để hạn chế drift.','Đây là biến thể nhóm kiểm chứng; KD là kỹ thuật đã có.'])
    slide('Thêm gì vào TPOUR?', ['Giữ CE + TRPO, reference momentum và queue.','Thêm KL trên xếp hạng + relative embedding MSE.','Teacher đóng băng; inference chỉ dùng student.','γ=10, η=1, temperature=0,1; cố định trước test.'])
    slide('Kiểm chứng công bằng', ['TRPO và anchored: seed 42/43/44, 1.000 bước mỗi nhánh.','Cùng dữ liệu, learning rate, batch và cách đánh giá.','Đủ 300 truy vấn SciFact; thêm ablation chỉ KL ở seed42.','Báo mean/SD, CI, số cập nhật thực và chi phí.'])
    slide('Kết quả toàn bộ SciFact',picture=root/'demo/01_full_test.png')
    slide('Demo: cùng query, hai mô hình',picture=root/'demo/03_ranking_demo.png')
    slide('Kết luận và giới hạn', [f'Mean nDCG@10: TRPO {groups["tpour"]["mean"]:.4f}; anchored {groups["anchored"]["mean"]:.4f}.',f'Chênh lệch với Contriever gốc: {baseline:+.4f}.','SciFact không đủ chứng minh temporal alignment.','Có cả ví dụ tốt/xấu, log thật và code chạy lại demo.'])
    pres.save(output/'Slide_cai_tien_TPOUR_retention.pptx')
    print(output)


if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--root',required=True); p.add_argument('--output',required=True)
    a=p.parse_args(); build(a.root,a.output)
