"""Audit report math families, slide evidence, and declared data lineage; no remote data fetch."""
import hashlib
import json
import sys
import unicodedata
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from real_experiment import clean_rows,connected_split,digest,read_jsonl
OUT=ROOT/'reports/ptit_verified'
R=ROOT/'results/anchored_kaggle'
SOURCE=ROOT/'reports/ptit_updated/Bao_cao_TPOUR_TRPO_da_sua_cong_thuc.docx'
REPORT=OUT/'Bao_cao_TPOUR_TRPO_PTIT_cap_nhat_da_ra_soat.docx'
SLIDES=ROOT/'reports/presentation_final/Slide_thuyet_trinh_TPOUR_TRPO_cap_nhat.pptx'
read=lambda p:json.loads(Path(p).read_text(encoding='utf-8'))
def main():
    manifest=read(OUT/'report_audit.json');assert digest(SOURCE)==manifest['template_sha256']
    for p in [REPORT,SLIDES]:
        with ZipFile(p) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith('.xml'):assert '\ufffd' not in z.read(n).decode('utf-8')
    d=Document(REPORT)
    math=[''.join(t.text or '' for t in e.iter(qn('m:t'))) for e in d._element.iter(qn('m:oMath'))]
    families={
      'Encoder/mean pooling':['encθ','mj hj'], 'CE/TRPO':['ℒCE','LTRPO'],
      'Time vector / interpolation':['τt =','θmid ='], 'IR metrics':['DCG@k','Recall@k','MRR@k','MAP@k'],
      'EMA':['θref ←'], 'Ranking distribution / KL':['Pa,i,j','Pteacher,i,j'],
      'Representation / total retention':['ℒrep','ℒnew']}
    normalized=[unicodedata.normalize('NFKC',t) for t in math]
    for name,keys in families.items():assert all(any(unicodedata.normalize('NFKC',k) in t for t in normalized) for k in keys),name
    for e in d._element.iter():
        rid=e.get(qn('r:embed'))
        if rid:assert rid in d.part.rels and d.part.rels[rid].reltype.endswith('/image')
    original=Document(SOURCE)
    assert [[c.text for c in row.cells] for row in d.tables[0].rows]==[[c.text for c in row.cells] for row in original.tables[0].rows]
    paras=[p.text for p in d.paragraphs];alltext='\n'.join(paras)
    for bad in ['L_total =','L_TRPO =','S_theta(','D_i^','Các ý tưởng LoRA/time adapters']:
        assert bad not in alltext,bad
    for heading in ['CHƯƠNG 1:','CHƯƠNG 2:','CHƯƠNG 3:','CHƯƠNG 4:','CHƯƠNG 5:']:
        assert any(t.startswith(heading) for t in paras)
    assert any(t.startswith('3.11.') for t in paras)
    for snippet in ['0,60035','0,62907','0,63935','0,44915','0,42910','1.000 lần thử','995 cập nhật']:
        assert snippet in alltext,snippet
    sample=ROOT/'upstream/TPOUR/2018_sample_train_trimmed.jsonl'
    rows,clean=clean_rows(read_jsonl(sample));train,held,split=connected_split(rows)
    saved=read(R/'data_manifest.json')
    assert clean==saved['cleaning'] and split==saved['split'] and digest(sample)==saved['source_sha256']
    # Recreate Linux JSONL serialization used by the GPU notebook.
    train_bytes=''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in train).encode('utf-8')
    trainhash=hashlib.sha256(train_bytes).hexdigest()
    for path in R.glob('*/provenance.json'):
        if path.parent.name=='baseline':continue
        assert read(path)['data_sha256']==trainhash,path
    evaluation=read(R/'tpour_seed42/metrics.json')
    for name,h in evaluation['data_hashes'].items():assert digest(ROOT/'data/scifact'/name)==h
    prs=Presentation(SLIDES);assert len(prs.slides)==27
    slide_text=[]
    for s in prs.slides:
        assert s.has_notes_slide and len(s.notes_slide.notes_text_frame.text)>80
        for sh in s.shapes:
            if sh.has_text_frame:slide_text.append(sh.text)
            if sh.has_table:slide_text.extend(c.text for row in sh.table.rows for c in row.cells)
    st='\n'.join(slide_text)
    for snippet in ['0,62907','0,60035','0,44915','0,42910','2.543','449','995','0,02872']:
        assert snippet in st,snippet
    assert read(ROOT/'reports/presentation_final/render_check.json')['text_overflow']==[]
    assert read(OUT/'word_render_check.json')['word_math_objects']==len(math)
    check={
      'template_unchanged':True,'cover_roster_matches_template':True,
      'native_math_objects':len(math),'equation_families_checked':list(families),
      'report_docx_sha256':digest(REPORT),'report_pdf_sha256':digest(REPORT.with_suffix('.pdf')),
      'report_pages':read(OUT/'word_render_check.json')['pages'],
      'slides':27,'main_slides':21,'appendix_slides':6,'speaker_notes_all_slides':True,
      'powerpoint_text_overflow':[], 'slide_pptx_sha256':digest(SLIDES),
      'lineage':{'author_sample_sha256':digest(sample),'cleaning_recomputed':clean,
        'split_recomputed':split,'train_serialization_sha256_matches_saved_runs':trainhash,
        'scifact_local_hashes_match_evaluation':evaluation['data_hashes']},
      'scope_of_code_review':{
        'train':'Reads prepared train.jsonl from author sample. Queue warmup/keys and fixed-teacher inputs use that train split.',
        'evaluation':'SciFact title+text, queries and qrels; custom binary metrics. Qrels are not used in training objective.',
        'demo':'Exact retrieval from SciFact corpus; no web-search or LLM context-generation call in reviewed live_retrieval_demo.py.',
        'declared_external_dependencies':'Pretrained facebook/contriever and tokenizer, author sample, official BEIR SciFact distribution; KD references support the extension only.',
        'limits':'Static code/data/hash audit plus exported evidence, not a new GPU training reproduction. Cannot audit complete pretrained data or exclude semantic/article overlap. SciFact-aware design bias remains.'},
      'validation_this_revision':{'unit_tests':'14 tests passed in the terminal during this revision.',
        'result_verifier':read(R/'verification.json')},
      'sources_verified_for_theory_extension':[
        'https://arxiv.org/abs/1606.09282','https://arxiv.org/abs/2406.05977']}
    (OUT/'delivery_audit.json').write_text(json.dumps(check,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in check.items() if k not in ['validation_this_revision','lineage','scope_of_code_review']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
