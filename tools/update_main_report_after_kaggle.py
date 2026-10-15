"""Revise the user's original Word report in place structurally, saving a separate updated file."""
import argparse
import copy
import hashlib
import json
import statistics
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
R=ROOT/'results/anchored_kaggle'


def read(path):return json.loads(Path(path).read_text(encoding='utf-8'))
def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def mn(tag,*children):
    e=OxmlElement('m:'+tag)
    for c in children:
        if isinstance(c,list):
            for x in c:e.append(x)
        else:e.append(c)
    return e
def mr(text):
    r=mn('r');p=OxmlElement('w:rPr');f=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+k),'Cambria Math')
    p.append(f);r.append(p);t=mn('t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);return r
def seq(*items):
    result=[]
    for item in items:
        if isinstance(item,list):result.extend(item)
        else:result.append(mr(item) if isinstance(item,str) else item)
    return result
def sub(b,s):return mn('sSub',mn('e',seq(b)),mn('sub',seq(s)))
def sup(b,s):return mn('sSup',mn('e',seq(b)),mn('sup',seq(s)))
def frac(a,b):return mn('f',mn('num',a if isinstance(a,list) else seq(a)),mn('den',b if isinstance(b,list) else seq(b)))
def loss(s):return sub('ℒ',s)


def build(source,output):
    before=sha(source);doc=Document(source);p=list(doc.paragraphs)
    if len(p)!=345 or not p[127].text.startswith('CHƯƠNG 3:'):raise ValueError('Unexpected template revision; inspect paragraph structure before editing')
    replace={
      3:'Thực nghiệm TPOUR trên dữ liệu thật, đánh giá SciFact và kiểm chứng biến thể giữ kiến thức của nhóm',
      5:'Hình 1. TPOUR và thành phần fixed-teacher retention bổ sung trong bản thực hành của nhóm.',
      10:'Báo cáo thuộc học phần Tìm kiếm và Truy xuất Dữ liệu, nghiên cứu vấn đề tài liệu đúng chủ đề nhưng không khớp thời gian của truy vấn. Nhóm triển khai một bản TPOUR/TRPO gọn bằng PyTorch, kiểm tra mã, thực nghiệm trên dữ liệu tác giả và đánh giá bằng SciFact. Sau các thử nghiệm ban đầu, nhóm bổ sung loss giữ kiến thức và kiểm chứng cả lợi ích lẫn đánh đổi của biến thể này.',
      11:'2. Mục đích: thực nghiệm có giới hạn và kiểm chứng biến thể mở rộng',
      12:'Mục tiêu là hiểu và triển khai các thành phần của bài báo đã chọn, thiết lập đối chứng có thể kiểm tra lại, đồng thời thử một biến thể retention. Đồ án chưa tái lập đầy đủ quy mô, temporal QA và nội suy time embedding của công trình gốc.',
      13:'Bài báo tham chiếu: HyunJin Kim, Jaejun Shim, Young Jin Kim và JinYeong Bak, “Temporal Preference Optimization for Unsupervised Retrieval”. Bản PDF 33 trang nhóm chọn ghi ICML 2026/PMLR 306; lưu tại upstream/TPOUR/ICML_2026_TPOUR.pdf. Young Jin Kim thuộc Microsoft; các tác giả còn lại thuộc Sungkyunkwan University theo bản PDF. Source tác giả: https://github.com/agwaBom/TPOUR [1].',
      14:'TPOUR chuyển ý tưởng học ưu tiên sang truy xuất có yếu tố thời gian. Temporal Retrieval Preference Optimization (TRPO) so sánh chênh lệch điểm giữa tài liệu ưu tiên và tài liệu lệch mốc, có mô hình tham chiếu cập nhật momentum và contrastive loss. TRPO trong báo cáo không phải Trust Region Policy Optimization của học tăng cường.',
      16:'Các phần việc đã có mã nguồn, log hoặc kết quả để kiểm tra gồm:',
      17:'Đọc bài báo được chọn; phân biệt phương pháp gốc, bản cài đặt gọn của nhóm và thành phần retention được thêm sau này.',
      18:'Triển khai Encoder, CE/TRPO, reference momentum và queue trong tpour.py; bổ sung runner dữ liệu thật, provenance, checkpoint/resume và đánh giá đủ SciFact trong real_experiment.py.',
      19:'Mở rộng từ 5 lên 14 kiểm tra tự động, gồm loss/gradient, queue, chia dữ liệu, metric, resume CPU và teacher đóng băng. Kết quả test xác nhận các trường hợp đã kiểm tra, không thay thế đánh giá chất lượng.',
      20:'Làm sạch mẫu Wikipedia tác giả: 3.000 dòng gốc còn 2.992 bộ ba; tách 2.543 train và 449 preference held-out. Train TRPO/retention ở ba seed trên Kaggle T4, cùng 1.000 lần thử bước mỗi run.',
      21:'Đánh giá đủ 300 truy vấn SciFact trên 5.183 tài liệu, tính lại metrics từ rankings/qrels; làm ablation và demo inference checkpoint thật. Ghi rõ retention cải thiện so với TRPO trên SciFact nhưng chưa vượt baseline và preference accuracy giảm.',
      25:'Bảng dưới là khung phân công của nhóm. Điền tên người thực hiện thực tế trước khi nộp; báo cáo không tự gán các hạng mục mới cho một cá nhân khi chưa có xác nhận.',
      68:'BM25 và các phương pháp thưa khai thác sự xuất hiện của từ trong tài liệu và truy vấn. Chúng có thể hạn chế khi hai văn bản diễn đạt cùng ý bằng từ khác nhau. Dense Retrieval dùng encoder để biểu diễn văn bản bằng vector; chất lượng vẫn phụ thuộc dữ liệu, mô hình và cách đánh giá, không mặc nhiên tốt hơn BM25 trong mọi trường hợp.',
      71:'Khi truy vấn có ngữ cảnh thời gian tường minh hoặc ngầm định, tài liệu đúng chủ đề có thể không khớp mốc cần tìm. Bảng dưới là ví dụ giả định về Tổ chức X để minh họa khái niệm, không phải dữ liệu thật hoặc bằng chứng chất lượng mô hình.',
      73:'Hai tài liệu cùng nói về Tổ chức X nhưng khác phiên bản quy trình. Một retriever chỉ dựa vào sự gần nhau ngữ nghĩa có thể đưa tài liệu lệch mốc lên trước; cần kiểm tra thêm sự phù hợp với thời gian của truy vấn.',
      75:'Mục tiêu là kiểm chứng bản thực hành dựa trên TPOUR và thử bổ sung retention để hạn chế suy giảm truy xuất tổng quát sau fine-tune. Các kết luận phải phân biệt kiểm thử mã, dữ liệu giả lập, benchmark tổng quát và đánh giá temporal alignment thực sự.',
      77:'Cài đặt bản gọn dựa trên phương pháp tác giả bằng PyTorch, hỗ trợ Toy Encoder để kiểm tra code và facebook/contriever cho dữ liệu thật; không gọi đây là chạy nguyên cấu hình distributed của tác giả.',
      78:'Kiểm tra công thức và chiều gradient bằng test tự động, đồng thời kiểm tra dữ liệu, resume và tính bất biến của teacher.',
      79:'Huấn luyện trên mẫu Wikipedia thật do tác giả phát hành, lưu cấu hình, số cập nhật thực, hash, phiên bản phần mềm, thời gian và bộ nhớ.',
      80:'Đánh giá đầy đủ SciFact để đo truy xuất tổng quát; bổ sung mined-preference proxy và phân tích các giới hạn trước khi kết luận về yếu tố thời gian.',
      81:'Hình 1 phân biệt student được tối ưu, reference cập nhật momentum và teacher cố định chỉ có ở biến thể retention. Mô hình inference chỉ dùng student; không trộn điểm Contriever gốc để tạo kết quả tốt hơn.',
      82:'Kết quả cần trả lời hai câu hỏi tách biệt: mô hình có giữ chất lượng truy xuất ngữ nghĩa không, và có ưu tiên đúng mốc thời gian không. SciFact chỉ hỗ trợ câu hỏi thứ nhất; chưa đủ bằng chứng để khẳng định mọi bài toán temporal retrieval đều cần hoặc được giải quyết bởi một kiến trúc duy nhất.',
      86:'Cơ sở phương pháp được đối chiếu với bài báo nhóm đã chọn [1]. Các phần retention và kiểm chứng mở rộng được ghi riêng là công việc của nhóm, không quy thành đóng góp sẵn có của tác giả.',
      87:'HyunJin Kim, Jaejun Shim, Young Jin Kim, JinYeong Bak. Temporal Preference Optimization for Unsupervised Retrieval. Bản PDF nhóm chọn ghi ICML 2026/PMLR 306, 33 trang. Source: https://github.com/agwaBom/TPOUR. Không dùng số liệu từ phiên bản arXiv khác để thay cho bản PDF này.',
      88:'Bài báo nghiên cứu temporal misalignment: tài liệu có thể gần ngữ nghĩa nhưng sai mốc truy vấn. Các thành phần nhóm tìm hiểu gồm học ưu tiên trên các phiên bản dữ liệu, TRPO, momentum reference, contrastive queue và nội suy time embedding. Bản thực hành hiện tại mới kiểm chứng một phần các thành phần đó.',
      89:'Học ưu tiên không giám sát: sử dụng cặp văn bản được khai thác từ các snapshot để tạo tín hiệu preferred/rejected, thay vì coi các cặp mined này là temporal gold do con người thẩm định.',
      90:'TRPO: so sánh chênh lệch similarity của hai tài liệu giữa model hiện tại và reference; dùng logistic loss lấy cảm hứng từ preference optimization.',
      91:'Momentum reference và queue cung cấp tham chiếu thay đổi chậm và mẫu âm cho contrastive learning. Cơ chế này không tự bảo đảm rằng fine-tune sẽ không làm giảm retrieval trên một benchmark khác.',
      92:'Các công thức dưới đây mô tả nền tảng phương pháp. Thành phần teacher cố định của nhóm được giới thiệu ở mục 2.6 và cài đặt tại Chương 3.',
      100:'Có thể lập chỉ mục vector tài liệu trước cho một encoder cố định. Khi đổi encoder hoặc trọng số nội suy, cần kiểm tra tính tương thích của index; chi phí phụ thuộc cách triển khai và khả năng tái sử dụng index.',
      104:'Công thức 3. Contrastive loss với nhiệt độ τ và negative embeddings từ queue.',
      107:'TRPO so sánh chênh lệch điểm ưu tiên tương đối với reference. Trong code nhóm, gap được chia nhiệt độ T trước khi nhân β. Ký hiệu ở Công thức 4a được viết tương ứng để tránh bỏ mất hệ số nhiệt độ.',
      113:'D⁺: tài liệu được bộ khai thác gán ưu tiên ở mốc mục tiêu; nhãn mined chưa được nhóm thẩm định độc lập.',
      114:'D⁻: tài liệu không được ưu tiên hoặc lệch mốc; có thể sớm hơn hoặc muộn hơn mốc truy vấn, không nhất thiết luôn cũ hơn.',
      115:'Reference dùng cập nhật EMA với m=0,9999 và không nhận gradient trực tiếp. Reference này khác teacher cố định bổ sung trong biến thể retention.',
      116:'β=0,3 điều khiển độ nhạy của preference loss trong cấu hình thực hành; T=0,05 là nhiệt độ similarity.',
      119:'Công thức 5. Objective CE+TRPO của bản thực hành; λ=0,925 không phải hệ số nội suy thời gian.',
      125:'MAP@k: trung bình average precision của từng truy vấn; tại mỗi hit, cộng precision ở vị trí đó, rồi chia tổng số tài liệu liên quan của truy vấn. Báo cáo dùng qrels nhị phân SciFact và giữ các truy vấn không có hit với điểm 0.',
    }
    for i,text in replace.items():
        props=copy.deepcopy(p[i].runs[0]._r.rPr) if i<8 and p[i].runs and p[i].runs[0]._r.rPr is not None else None
        p[i].clear();r=p[i].add_run(text)
        if props is not None:r._r.insert(0,props)
    # Replace the cover picture with the code-generated method figure.
    p[4].clear();p[4].add_run().add_picture(str(ROOT/'reports/anchored_improvement/method.png'),width=Inches(6.1))
    # Correct temperature scaling in the preserved native equation.
    p[108].clear();equation=mn('oMath')
    for item in seq(sub('Δ','θ'),' = ',frac(seq(sub('s','θ'),'(Q, D⁺) − ',sub('s','θ'),'(Q, D⁻)'),'T'),',  ',sub('Δ','ref'),' = ',frac(seq(sub('s','ref'),'(Q, D⁺) − ',sub('s','ref'),'(Q, D⁻)'),'T')):equation.append(item)
    p[108]._p.append(equation)
    p[94].clear();p[94].add_run('BM25 là hàm xếp hạng từ vựng dựa trên TF, IDF và chuẩn hóa độ dài tài liệu. Nó được giới thiệu ở đây để đối chiếu với biểu diễn dense, không đồng nhất hoàn toàn với mô hình vector cổ điển.')
    # Preserve assignment names/placeholders, update only the responsibility column.
    duties=[
      'Tổng hợp phạm vi nghiên cứu, quản lý mã nguồn và bằng chứng; rà soát kết luận, giới hạn và báo cáo cập nhật.',
      'Đối chiếu bài báo TPOUR, CE/TRPO và các nguồn knowledge distillation; phân biệt phương pháp gốc với retention nhóm thử.',
      'Triển khai Encoder, queue, EMA, runner dữ liệu thật và teacher cố định/retention loss.',
      'Kiểm thử mã, chuẩn bị dữ liệu, chạy đối chứng ba seed trên Kaggle; kiểm tra checkpoint, số bước và provenance.',
      'Đánh giá SciFact, bootstrap/ablation, demo inference thực; chuẩn bị biểu đồ, slide và đối chiếu giới hạn.'
    ]
    for row,text in zip(doc.tables[0].rows[1:],duties):row.cells[2].text=text
    # Remove old manual TOC entries and the old chapters 3 onward, including stale tables/pictures.
    for x in p[28:64]:x._p.getparent().remove(x._p)
    start=p[127]._p;body=doc._element.body;remove=False
    for child in list(body):
        if child is start:remove=True
        if remove and child.tag!=qn('w:sectPr'):body.remove(child)
    def para(text,style=None):return doc.add_paragraph(text,style)
    def heading(text,level=2):
        x=doc.add_heading(text,level)
        if level==1:x.paragraph_format.page_break_before=True
        return x
    def eq(parts,caption):
        x=para('');m=mn('oMath')
        for item in parts:m.append(item)
        x._p.append(m);para(caption,'Caption')
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
        pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
        for row in rows:
            for c,text in zip(t.add_row().cells,row):c.text=str(text)
        for row in t.rows:
            for c in row.cells:
                for x in c.paragraphs:
                    for run in x.runs:run.font.name='Times New Roman';run.font.size=Pt(10.5)
        para('')
        return t
    def figure(path,caption):
        para('').add_run().add_picture(str(path),width=Inches(6.2));para(caption,'Caption')
    # Insert a concise updated abstract after the cover, before the old introduction.
    marker=p[8]._p
    for text,style in [
      ('TÓM TẮT BẢN CẬP NHẬT', 'Heading 1'),
      ('Báo cáo nghiên cứu TPOUR dựa trên bài báo đã chọn [1], triển khai CE/TRPO và mở rộng bằng fixed-teacher retention. Sau demo giả lập và smoke test, nhóm chuyển sang Wikipedia thật, làm sạch/tách dữ liệu, đánh giá đủ SciFact và chạy đối chứng ba seed trên Kaggle. Phần kiểm chứng hiện có mã, cấu hình, log, rankings, hash và demo checkpoint thật.', 'Normal'),
      ('Trên 300 truy vấn SciFact, TRPO đối chứng đạt nDCG@10 trung bình 0,60035 ± 0,00432; biến thể retention đạt 0,62907 ± 0,00204. Contriever gốc đạt 0,63935. Preference accuracy giữ riêng giảm từ 0,44915 xuống 0,42910. Do đó, kết quả chỉ hỗ trợ việc hạn chế suy giảm truy xuất tổng quát so với TRPO trong cấu hình này; chưa chứng minh cải thiện toàn diện hoặc temporal alignment.', 'Normal'),
      ('Số liệu nDCG@10=0,74472 của giai đoạn thăm dò chỉ thuộc 30 truy vấn, không đại diện kết quả đủ 300 truy vấn. Demo giả lập chỉ kiểm tra pipeline. Các bằng chứng hiện có chưa đủ để khẳng định không quên tri thức, tái lập toàn diện bài báo hoặc cải thiện temporal retrieval.', 'Normal')]:
        x=doc.add_paragraph(text,style);marker.addprevious(x._p)
    # Correct front heading and create a manual outline without invented page numbers.
    p[27].clear();p[27].add_run('Mục lục nội dung (theo cấu trúc bản cập nhật)')
    toc=[]
    for text in [
      'TÓM TẮT BẢN CẬP NHẬT; LỜI MỞ ĐẦU VÀ PHẠM VI THỰC HIỆN',
      'CHƯƠNG 1. GIỚI THIỆU VÀ PHÁT BIỂU BÀI TOÁN',
      'CHƯƠNG 2. CƠ SỞ LÝ THUYẾT TỪ TPOUR VÀ RETENTION',
      '  2.1–2.5. Bài báo gốc, dense retrieval, CE/TRPO và metrics',
      '  2.6. Fixed-teacher knowledge retention và giới hạn kỳ vọng',
      'CHƯƠNG 3. THIẾT KẾ, DỮ LIỆU VÀ CÀI ĐẶT CẬP NHẬT',
      '  3.1. Nguồn dữ liệu; 3.2. Làm sạch và tách dữ liệu',
      '  3.3. Encoder/reference/teacher; 3.4. Objective và thành phần mới',
      '  3.5. Train/resume; 3.6. Đánh giá; 3.7. Mã nguồn và demo',
      '  3.8. Lập luận thiết kế; 3.9. Thuật toán; 3.10. Đóng góp và tái lập',
      'CHƯƠNG 4. THỰC NGHIỆM, KẾT QUẢ VÀ ĐÁNH ĐỔI',
      '  4.1. Các giai đoạn; 4.2. Đối chứng; 4.3. Kết quả ba seed',
      '  4.4. Bootstrap; 4.5. Preference; 4.6. Ablation MSE',
      '  4.7. Chi phí; 4.8. Demo checkpoint thật; 4.9. Giới hạn',
      'CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN',
      '  5.1–5.3. Công việc, hướng mở và đối chiếu mục tiêu',
      '  5.4–5.6. Ý nghĩa, rào cản và kế hoạch kiểm chứng',
      '  5.7. Các câu hỏi trọng tâm khi bảo vệ đồ án',
      'PHỤ LỤC. LỆNH CHẠY, ĐẦU RA, ĐỐI CHIẾU VÀ TÀI LIỆU THAM KHẢO']:
        x=doc.add_paragraph(text);toc.append(x)
    marker=p[65]._p
    for x in toc:marker.addprevious(x._p)
    heading('2.6. Fixed-teacher knowledge retention và giới hạn kỳ vọng')
    para('Knowledge distillation có thể dùng để giữ phản hồi của model khởi tạo trong khi fine-tune [6, 7]. Trong biến thể của nhóm, teacher là bản Contriever ban đầu đóng băng; reference của TRPO vẫn cập nhật EMA. Hai vai trò này khác nhau: reference tham gia preference gap/queue, teacher cung cấp mục tiêu giữ phân bố xếp hạng và biểu diễn.')
    para('Giả thuyết là hạn chế drift có thể giảm mức suy giảm retrieval sau fine-tune trên mẫu nhỏ. Teacher cũng có thể bảo lưu thông tin lệch thời gian hoặc hạn chế cập nhật mới. Vì vậy cần đo cả SciFact và preference proxy, không mặc nhiên coi retention là cải thiện temporal retrieval. Nhóm không tuyên bố phát minh knowledge distillation hoặc tính mới đã được peer review.')
    heading('CHƯƠNG 3: THIẾT KẾ, DỮ LIỆU VÀ CÀI ĐẶT CẬP NHẬT',1)
    heading('3.1. Nguồn dữ liệu và phạm vi kế thừa')
    para('Nguồn train là upstream/TPOUR/2018_sample_train_trimmed.jsonl, gồm 3.000 dòng từ tác giả. SHA-256: '+read(R/'data_manifest.json')['source_sha256']+'. Nhóm dùng mẫu đã phát hành; chưa tự thu thập lại toàn bộ snapshot, khai phá FAISS và tái lập corpus nhiều thời điểm của tác giả. Không gán thời gian mới chỉ dựa vào các con số trong nội dung.')
    para('SciFact thuộc BEIR [3], gồm corpus 5.183 tài liệu, queries và qrels. Lần chạy mới dùng đủ 300 truy vấn test có nhãn, không giới hạn ở 30 truy vấn đầu như số liệu lịch sử. SciFact đo tìm tài liệu liên quan cho khẳng định khoa học, chưa phải benchmark kiểm chứng mốc thời gian.')
    heading('3.2. Làm sạch, chia dữ liệu và các giới hạn')
    table(['Hạng mục','Số lượng','Ý nghĩa'],[
      ['Dòng gốc',3000,'Mẫu Wikipedia từ tác giả'],['Thiếu context bị loại',4,'Không đủ preferred/rejected'],['Cặp giống nhau bị loại',4,'Preferred và rejected giống sau chuẩn hóa'],['Bộ ba hợp lệ',2992,'Không phải dữ liệu mock'],['Train',2543,'Dùng để cập nhật model'],['Held-out preference',449,'Không dùng để train; nhãn mined, không phải temporal gold']])
    para('Các bộ ba có cùng query hoặc tài liệu sau chuẩn hóa khoảng trắng/chữ hoa được nối thành cùng nhóm rồi chia bằng seed42. Mọi seed huấn luyện dùng cùng split này. Giao exact text giữa train/held-out bằng 0; chưa loại mọi trùng article, excerpt hoặc paraphrase vì không có article ID đáng tin cậy. Không gọi đây là bảo đảm loại hoàn toàn data leakage.')
    para('Input gốc có question, positive_ctxs, weak_positive_ctxs; context có thể là chuỗi hoặc object chứa text. Bản gọn lấy context đầu tiên mỗi loại, chuẩn hóa thành query/preferred/rejected và lưu source_line. Tokenizer giới hạn 256 token. Audit train trước augmentation thấy 4/2.543 cặp trở thành giống nhau sau truncation; word deletion có thể làm thay đổi các token xuất hiện. Dữ liệu không được gán thêm timestamp giả.')
    heading('3.3. Ba vai trò encoder và inference')
    table(['Thành phần','Cập nhật','Vai trò'],[
      ['Current/student','AdamW, nhận gradient','Mã hóa query và tài liệu, tối ưu objective'],['Momentum reference','EMA m=0,9999; no_grad/eval','Preference reference và document keys trong queue'],['Fixed teacher (chỉ retention)','Đóng băng từ lúc khởi tạo; no_grad/eval','Mục tiêu distillation, không cập nhật EMA']])
    para('Bi-encoder mô tả vai trò mã hóa query/tài liệu; trong bản này chúng chia sẻ trọng số current encoder. Không nhầm hai vai trò query/document với hai mạng current/reference. Khi inference chỉ tải student encoder. Teacher, reference và queue không được trộn điểm vào xếp hạng. Toy Encoder 4.096×64 tham số embedding chỉ dùng kiểm tra mã; Contriever dùng Transformer và vector 768 chiều.')
    eq(seq(sub('θ','ref'),' ← m ',sub('θ','ref'),' + (1 − m) ',sub('θ','student')), 'Công thức 6. Cập nhật reference EMA; teacher cố định không áp dụng phép cập nhật này.')
    heading('3.4. Objective bổ sung của nhóm: ranking KL và relative MSE')
    para('Candidate list của mỗi batch gồm preferred và rejected của chính batch đó. Student và teacher nhìn cùng các chuỗi đã augmentation. Tạo phân bố softmax từ cosine similarity chia nhiệt độ retention 0,1; dùng KL(teacher || student), trung bình theo query. Teacher targets được detach. Đây là distillation trên candidate list, không phải reward model hoặc nhãn gold mới.')
    eq(seq(loss('KL'),' = KL(',sub('P','teacher'),' ∥ ',sub('P','student'),')'), 'Công thức 7. Ranking retention trên cùng candidate list; nhiệt độ 0,1.')
    para('Relative representation loss là trung bình squared L2 error giữa embedding student/teacher, chia squared teacher norm với chặn dưới ε=10⁻⁶; lấy trên query/preferred/rejected. Thành phần này giữ cả hướng lẫn độ lớn vector, vì inference dùng dot product không normalize. Retention được tính FP32 kể cả khi train AMP.')
    eq(seq(loss('rep'),' = mean ',frac(seq(sup(seq('‖',sub('z','student'),' − ',sub('z','teacher'),'‖'),'2')),seq('max(',sup(seq('‖',sub('z','teacher'),'‖'),'2'),', ε)'))), 'Công thức 8. Relative embedding MSE, ε=10⁻⁶.')
    eq(seq(loss('new'),' = λ ',loss('CE'),' + (1 − λ) ',loss('TRPO'),' + γ (',loss('KL'),' + η ',loss('rep'),')'), 'Công thức 9. Objective biến thể retention; λ=0,925, γ=10, η=1.')
    para('Bản gốc thực hành chỉ dùng hai hạng đầu. Biến thể giữ nguyên CE/TRPO/reference/queue và thêm hạng retention. Hệ số γ=10, η=1, temperature retention=0,1 được cố định trước lần chạy mới; không chọn checkpoint tốt nhất theo SciFact. Ablation η=0 ở seed42 kiểm tra phần MSE. Kế hoạch và code trước chạy được lưu ở commit 219a21a.')
    heading('3.5. Train, ổn định số học và checkpoint/resume')
    para('Runner nạp vector thật vào queue trước vòng tối ưu, lấy batch ngẫu nhiên, xóa từ 10%, forward current/reference và teacher nếu có, tính loss, backward AMP, unscale và clip norm 1,0. Nếu gradient không hữu hạn, scaler bỏ cập nhật; chỉ cập nhật scheduler, EMA và queue khi bước tối ưu hợp lệ. Mọi run mới có 1.000 lần thử bước, 995 cập nhật thực. Đây không phải tuyên bố mỗi mẫu đã được quét đủ một epoch.')
    para('Checkpoint train lưu current/reference/teacher (nếu có), queue, optimizer, scheduler, scaler và RNG; encoder.pt chỉ lưu student để inference. Resume yêu cầu cấu hình/dữ liệu giống lần đầu. Tính tương đương resume đã kiểm tra trên bài CPU nhỏ; chưa khẳng định bitwise reproducibility trên mọi GPU. Trong notebook, sau train và hai phép evaluate thành công, training_state lớn được bỏ để tiết kiệm đĩa; encoder.pt/log/config vẫn giữ ở Kaggle Output.')
    para('Queue dùng O(Kd) bộ nhớ theo số keys K và chiều vector d, không phải O(1) khi tăng K hoặc d. Với K cố định, queue tái sử dụng buffer vòng qua các bước. Warmup bảo đảm có negative keys từ đầu; không tự bảo đảm mọi cấu hình đều hội tụ hoặc không có NaN.')
    heading('3.6. Đánh giá, cache và provenance')
    para('Mọi nhánh dùng model facebook/contriever, max_length256, dot product, không normalize. Corpus được ghép title + text và mã hóa theo cùng thứ tự ID; xếp hạng exact bằng score, tie-break ổn định. Lưu top-100, metrics từng query và trung bình, kể cả query không có hit. Chỉ số được tính bằng code đã kiểm tra với ví dụ tay, rồi tính lại từ run.json và qrels; không gọi đây là trực tiếp chạy thư viện beir/pytrec_eval nếu runner không sử dụng chúng.')
    para('Lưu SHA-256 của corpus/queries/qrels/checkpoint, commit và hash mã, phiên bản phần mềm, GPU và cấu hình. Index cache của live demo kiểm tra hash model/corpus, ID, max length, normalization và revision trước khi dùng lại. Khi đổi encoder phải kiểm tra hoặc xây lại index; báo cáo chưa đo hệ thống quy mô triệu tài liệu hoặc khẳng định độ trễ micro-giây.')
    heading('3.7. Cấu trúc mã nguồn và phần đã bổ sung')
    table(['File','Vai trò'],[
      ['tpour.py','Encoder, CE/TRPO, EMA, queue; giữ demo toy để smoke test'],['real_experiment.py','Prepare/train/resume/full evaluate/preference/bootstrap/provenance'],['anchored_tpour.py','Teacher cố định, ranking KL và relative embedding MSE'],['test_tpour.py / test_real_experiment.py / test_anchored_tpour.py','5 + 5 + 4 = 14 kiểm tra'],['verify_anchor_results.py','Đối chiếu protocol/config/hash, tính lại metrics/CI và live evidence'],['demo_retention.py','Replay rankings của run thật, ghi rõ chế độ'],['live_retrieval_demo.py','Inference encoder thật với query nhập và cache index'],['notebooks/TPOUR_anchored_Kaggle.ipynb','Đối chứng ba seed và ablation'],['notebooks/TPOUR_live_checkpoint_demo.ipynb','Demo từ checkpoint output và truy vấn mới']])
    heading('3.8. Lập luận thiết kế: từ vấn đề quan sát đến giả thuyết kiểm chứng')
    para('Thiết kế retention xuất phát từ việc fine-tune trên mẫu Wikipedia nhỏ làm giảm SciFact so với Contriever khởi tạo. Đây là dấu hiệu suy giảm trên một benchmark tổng quát; chưa đủ để kết luận cơ chế duy nhất là catastrophic forgetting. Thay đổi miền dữ liệu, nhãn mined và cấu hình tối ưu đều có thể góp phần. Nhóm đặt giả thuyết hẹp: ràng buộc student gần hành vi ban đầu có thể giữ được một phần năng lực truy xuất tổng quát trong khi vẫn học CE/TRPO.')
    table(['Quyết định','Lý do và hệ quả cần kiểm tra'],[
      ['Teacher cố định, tách EMA reference','EMA đi theo student nên không còn là mốc hành vi ban đầu. Teacher cố định tạo mốc giữ kiến thức, nhưng cũng có thể cản học thông tin mới.'],
      ['KL trên thứ tự tương đối của batch','Giữ phân bố ưu tiên mềm thay vì chỉ ép một tài liệu đúng. Candidate set nhỏ và phụ thuộc batch, chưa đại diện toàn corpus.'],
      ['Cosine cho distillation, dot cho retrieval','Cosine giảm nhạy với scale khi tạo phân bố; dot giữ protocol Contriever. Hai không đồng nhất, nên cần đo retrieval thực và thử thành phần giữ biểu diễn.'],
      ['Relative MSE bổ sung','Ràng buộc cả hướng và norm; chia norm teacher giúp so lỗi tương đối. Hiệu quả riêng phải kiểm tra bằng ablation, không suy ra từ ý nghĩa công thức.'],
      ['Inference chỉ dùng student','Có thể dùng cùng quy trình lập chỉ mục như đối chứng; phần teacher làm tăng chi phí train, không thêm một lượt teacher khi trả lời query.']])
    para('Từ giả thuyết này, thí nghiệm phải có ba góc nhìn: SciFact đo retrieval tổng quát; held-out preference theo dõi mục tiêu mined; ablation phân biệt tác động của KL và MSE. Nếu chỉ SciFact tăng mà preference giảm, kết quả được diễn giải là đánh đổi giữa hai mục tiêu. Thiết kế không mặc định retention sẽ tốt ở mọi mốc hoặc mọi miền dữ liệu.')
    heading('3.9. Thuật toán một bước huấn luyện và luồng suy luận')
    for text in [
      'Bước 1 — Đọc batch bộ ba từ train split cố định, tạo augmentation và tokenize tối đa 256 token. Cùng văn bản đã augmentation được đưa cho student và teacher để tránh distillation hai đầu vào khác nhau.',
      'Bước 2 — Student tạo embedding query, preferred và rejected có gradient. Reference tạo biểu diễn phục vụ preference gap và keys không gradient; teacher tạo mục tiêu distillation không gradient.',
      'Bước 3 — Tính CE bằng positive và negative queue; tính TRPO theo chênh lệch student/reference. Với biến thể mới, ghép 2B tài liệu preferred/rejected của batch cho mỗi query để tính KL; tính relative MSE trên cả ba nhóm embedding.',
      'Bước 4 — Cộng objective theo Công thức 9, backward bằng AMP, unscale gradient, kiểm tra hữu hạn và clip. Nếu bước không hợp lệ thì bỏ cập nhật; nếu hợp lệ mới cập nhật optimizer, scheduler, EMA reference và queue.',
      'Bước 5 — Lưu log, cấu hình và checkpoint. Đánh giá với model cuối theo protocol cố định; không duyệt các checkpoint theo điểm test để chọn kết quả cao nhất.'
    ]:para(text)
    para('Ví dụ batch size B=4 tạo 8 candidate documents cho mỗi query khi distillation, tức ma trận score 4×8. Đây là các mẫu trong batch, không phải 8 tài liệu do qrels SciFact chọn. Qrels test chỉ được dùng ở bước đánh giá; không dùng làm nhãn train. Reference keys đưa vào queue và candidate list distillation có vai trò khác nhau, không nên đồng nhất hai cấu trúc.')
    para('Luồng inference: tải encoder.pt của student → kiểm tra hash/protocol cache → mã hóa corpus hoặc dùng index tương thích → mã hóa query mới → tính dot product với toàn corpus → lấy top-k và lưu JSON. Trong demo replay, thay bước mã hóa bằng việc đọc run.json đã đánh giá; hai chế độ được ghi rõ để tránh trình bày kết quả lưu sẵn như một lần chạy mới.')
    heading('3.10. Phân định đóng góp và khả năng tái lập')
    table(['Lớp công việc','Đã thực hiện','Phạm vi chưa bao phủ'],[
      ['Kế thừa công trình gốc','Contriever, ý tưởng temporal preference, CE/TRPO, reference và queue','Chưa tái lập toàn bộ dữ liệu, lịch train và temporal QA của tác giả'],
      ['Kỹ thuật thực nghiệm của nhóm','Làm sạch/split, runner, provenance, kiểm thử, metrics và demo checkpoint','Chưa bảo đảm loại trùng ngữ nghĩa hoặc tái lập bitwise trên mọi phần cứng'],
      ['Biến thể phương pháp của nhóm','Teacher cố định và objective KL + relative MSE, có ablation','Chưa chứng minh tính mới học thuật; MSE chưa có lợi ích riêng chắc chắn'],
      ['Bằng chứng đánh giá','Full SciFact, ba seed, bootstrap query-level, proxy và chi phí','Chưa có benchmark độc lập chưa dùng khi thiết kế hoặc temporal gold']])
    para('Điều kiện tối thiểu để người khác kiểm tra lại là cùng dữ liệu/hash, cùng split và tokenizer/revision, cùng cấu hình objective, số cập nhật thực và quy tắc metric. Có thể kiểm tra rankings/metrics từ GitHub mà không train lại; muốn inference query mới cần tải đúng checkpoint từ Kaggle. Muốn tái huấn luyện cần GPU phù hợp và cài môi trường, không thể chỉ mở ảnh biểu đồ rồi coi là đã tái lập.')
    heading('CHƯƠNG 4: THỰC NGHIỆM, KẾT QUẢ VÀ ĐÁNH ĐỔI',1)
    heading('4.1. Các giai đoạn thực hiện và giá trị bằng chứng')
    table(['Giai đoạn','Đã chạy','Cách diễn giải'],[
      ['Demo toy ban đầu','32 bộ ba train, 8 query, 16 docs; 150 bước; CE và TRPO đều đạt 1,0','Chỉ kiểm tra pipeline; mẫu câu thuận lợi, chưa chứng minh TRPO hơn CE'],['Smoke/phiên Hạnh','Contriever 2 bước; sau đó log 20 bước và SciFact 30 query','Số 0,74472 chỉ thuộc subset; không so trực tiếp với full300 hoặc kết luận không quên'],['Kaggle dữ liệu thật lần đầu','CE-only và CE+TRPO, seed42, 1.000 lần thử bước; full300','CE nDCG@10=0,53975; TRPO=0,59560; baseline=0,63935. Log ở results/kaggle_real_seed42'],['Kaggle retention mới','Ba seed TRPO/retention; ablation chỉ KL seed42; full300','Bằng chứng chính của bản cập nhật; source train 219a21a'],['Demo checkpoint thật','Bốn model replay query cố định và một query mới','Kiểm tra inference/trace, không thay thế metrics toàn test']])
    para('Các giai đoạn dùng dữ liệu hoặc protocol khác nhau phải trình bày riêng. Không lấy toy query37, số liệu 30 query và kết quả đủ 300 query ghép thành một bằng chứng. Loss thấp hơn giữa hai objective khác nhau không đủ kết luận chất lượng retrieval tốt hơn. Lần seed42 mới hơi khác lần đầu do điều kiện thực thi/numerics; bảng đối chứng chính dùng các run mới cùng protocol.')
    heading('4.2. Thiết lập đối chứng và câu hỏi nghiên cứu')
    cfg=read(R/'tpour_seed42/config.json');env=read(R/'tpour_seed42/provenance.json')['environment']
    table(['Hạng mục','Giá trị'],[
      ['Thiết bị','Tesla T4; mỗi run dùng một GPU, hai nhánh có thể chạy song song trên T4×2'],['Phần mềm',f'Python {env["python"]}; PyTorch {env["packages"]["torch"]}; Transformers {env["packages"]["transformers"]}'],['Train/held-out','2.543 / 449; split cố định seed42'],['Seed train','42, 43, 44 cho TRPO và retention'],['Steps / batch / queue','1.000 lần thử / 4 / 4.096'],['Optimizer','AdamW lr1e-6, betas=(0,9; 0,98), epsilon1e-6; warmup100 + linear decay'],['CE/TRPO','λ=0,925; β=0,3; T=0,05; EMA m=0,9999'],['Augmentation / precision','Word deletion0,1; AMP; gradient checkpointing; clipping1,0'],['Retention','γ=10; η=1; temperature0,1; ablation η=0 seed42'],['Evaluation','5.183 docs / 300 queries; max_length256; dot; không normalize; top100']])
    para('RQ1: triển khai có vượt các test code và có truy nguyên dữ liệu/checkpoint không? RQ2: retention có giảm mức suy giảm SciFact so với TRPO không? RQ3: mined-preference proxy và chi phí thay đổi thế nào? RQ4: MSE có lợi riêng so với chỉ KL không? Không đặt sẵn câu trả lời “đạt xuất sắc”; kết luận theo số liệu.')
    heading('4.3. Kết quả chính trên đủ 300 truy vấn SciFact')
    groups=read(R/'demo/summary.json')
    table(['Nhánh','nDCG@10 mean','Sample SD','Số run'],[[n,f'{groups[k]["mean"]:.5f}',f'{groups[k]["sample_sd"]:.5f}' if groups[k]['sample_sd'] is not None else '—',groups[k]['n']] for n,k in [('Contriever gốc','baseline'),('TRPO đối chứng','tpour'),('Chỉ thêm ranking KL','ranking_only'),('KL + relative MSE','anchored')]])
    para('Retention cao hơn TRPO trung bình 0,02872 nDCG@10, tương đương 2,872 điểm trên thang100 hoặc khoảng 4,78% tương đối. Tuy nhiên retention vẫn thấp hơn Contriever gốc trung bình 0,01028. SD là sample SD qua ba seed, không phải CI tổng quát cho mọi seed/domain. Baseline là model cố định; ablation chỉ có một seed.')
    figure(R/'demo/01_full_test.png','Hình 2. Kết quả toàn bộ SciFact; error bars là SD giữa các seed. Không dùng số liệu 30 query cũ thay cho bảng này.')
    heading('4.4. Chênh lệch ghép cặp và khoảng tin cậy')
    rows=[]
    for seed in [42,43,44]:
        c=read(R/f'compare_tpour_seed{seed}_anchored_seed{seed}.json')
        rows.append([seed,f'{read(R/f"tpour_seed{seed}/metrics.json")["metrics"]["NDCG@10"]:.5f}',f'{read(R/f"anchored_seed{seed}/metrics.json")["metrics"]["NDCG@10"]:.5f}',f'{c["mean_delta"]:+.5f}',f'[{c["ci95"][0]:+.5f}; {c["ci95"][1]:+.5f}]'])
    table(['Seed','TRPO','Retention','Delta','CI95% delta'],rows)
    para('Bootstrap ghép cặp theo đúng query ID, 5.000 lần, seed123. CI retention−TRPO dương ở cả ba seed trong phép đo này. CI retention−baseline chứa0 ở seed42/43 và âm ở seed44; chưa có bằng chứng vượt Contriever gốc. CI theo query của từng run không thay thế bất định giữa seed hoặc giữa benchmark.')
    heading('4.5. Preference proxy và mức giữ xếp hạng teacher')
    rows=[]
    for seed in [42,43,44]:
        a=read(R/f'tpour_seed{seed}/heldout_preference.json')['pairwise_preference_accuracy'];b=read(R/f'anchored_seed{seed}/heldout_preference.json')['pairwise_preference_accuracy']
        rows.append([seed,f'{a:.5f}',f'{b:.5f}',f'{b-a:+.5f}'])
    table(['Seed','TRPO accuracy','Retention accuracy','Delta'],rows)
    para('Accuracy mined preference trung bình giảm từ 0,44915 xuống 0,42910; cả hai dưới0,5. Đây là đánh đổi quan sát được, không được bỏ khỏi kết luận. Proxy chỉ đo thứ tự score đối với nhãn khai thác từ mẫu tác giả, chưa phải temporal QA gold được nhóm thẩm định độc lập. Không suy ra chất lượng temporal alignment chỉ từ SciFact tăng.')
    para('Top-10 overlap với baseline tăng trung bình từ khoảng0,7173 lên0,8884, phù hợp mục tiêu giữ xếp hạng teacher. Agreement này không phải accuracy và không chứng minh giữ đúng thông tin theo thời gian; teacher có thể duy trì cả các ưu tiên chưa phù hợp mốc.')
    heading('4.6. Ablation: phần relative MSE có lợi riêng không?')
    c=read(R/'compare_ranking_only_full_seed42.json')
    table(['Seed42','nDCG@10'],[['Chỉ ranking KL','0,62870'],['KL + relative MSE','0,63082'],['Delta / CI95%',f'{c["mean_delta"]:+.5f} / [{c["ci95"][0]:+.5f}; {c["ci95"][1]:+.5f}]']])
    para('CI chứa0, nên chưa đủ bằng chứng cho lợi ích riêng của MSE về nDCG. Không được suy từ kết quả tổng thể rằng mọi thành phần đều cần thiết. Cần lặp ablation nhiều seed và benchmark chưa dùng khi thiết kế; chưa có phép tìm hyperparameter tối ưu hoặc tuyên bố biến thể tốt nhất.')
    heading('4.7. Số cập nhật, chi phí và diễn biến loss')
    table(['Run','Cập nhật thực','Train (s)','Peak allocated GiB'],[[x.parent.name,read(x)['optimizer_updates'],f'{read(x)["total_seconds"]:.1f}',f'{read(x)["peak_gpu_bytes"]/2**30:.2f}'] for x in sorted(R.glob('*/completed.json'))])
    para('Mỗi run có 995 cập nhật hợp lệ/1.000 lần thử; AMP bỏ5 bước do gradient không hữu hạn. Chi phí train trung bình khoảng332,2 giây cho TRPO và385,3 giây cho retention; peak allocated khoảng2,51 và2,92 GiB. Số thời gian không gồm tải model, queue warmup và evaluate; peak allocated không phải tổng VRAM thiết bị.')
    figure(R/'demo/02_training.png','Hình 3. Log loss thực; độ lớn các objective khác nhau không dùng để xếp hạng chất lượng mô hình.')
    heading('4.8. Demo và kiểm chứng checkpoint thật')
    para('14 kiểm tra tự động đã qua ở máy và trên Kaggle: 5 test cơ bản, 5 test runner/dữ liệu/metrics/resume, 4 test retention. Teacher được kiểm tra không nhận gradient/không đổi khi optimizer và EMA chạy; gamma0 cho gradient giống TRPO; loss bằng0 khi embeddings giống nhau và giảm theo hướng gradient. Các test là kiểm tra những trường hợp cụ thể, không phải chứng minh đúng tuyệt đối mọi đầu vào.')
    checks=read(R/'live_demo/verification.json');err=max(c['max_score_error'] for c in checks)
    para(f'Job live inference đã COMPLETE cho Contriever gốc, TRPO, ranking-only và retention seed42. Cả bốn khớp top3 query1 với rankings đã lưu; sai số score tối đa {err:.2e}. Hash checkpoint thật tính trong job inference khớp bản train/evaluation. Sau đó model retention chạy query mới “What influences the mechanical properties of biological materials?” bằng index thực. Query mới chưa có gold labels nên không báo accuracy.')
    figure(R/'demo/06_live_inference.png','Hình 4. Render từ JSON inference thực của checkpoint; không phải ảnh AI hoặc ảnh giao diện web giả.')
    para('Có hai chế độ demo: replay run.json để trình bày benchmark đã đo, và live_retrieval_demo.py để mã hóa query nhập mới bằng encoder/checkpoint. Không gọi replay là inference mới. Query1 chọn cố định trước và có trường hợp không tìm đúng trong top3; cần trình bày cả giới hạn này.')
    figure(R/'demo/04_ranking_demo.png','Hình 5. Query1282 được chọn sau đánh giá vì cải thiện lớn nhất ở seed42; tài liệu gold từ hạng3 lên hạng1. Baseline gốc cũng tìm được gold ở hạng1.')
    figure(R/'demo/05_ranking_demo.png','Hình 6. Query314 được chọn vì giảm nhiều nhất ở seed42; minh họa biến thể vẫn có trường hợp kém TRPO. Ví dụ được chọn không thay thế thống kê toàn test.')
    heading('4.9. Mức độ kiểm chứng và giới hạn')
    table(['Câu hỏi','Bằng chứng','Kết luận phù hợp'],[
      ['RQ1: code/provenance','14 test; hash/config/metrics đối chiếu; live replay khớp','Đạt các kiểm tra đã thực hiện, chưa phải tái lập toàn bộ bài báo'],['RQ2: SciFact so với TRPO','Mean +0,02872; cả ba CI query-level dương','Hạn chế suy giảm so với TRPO trong cấu hình này; chưa vượt baseline'],['RQ3: proxy/chi phí','Preference accuracy giảm; train/memory tăng','Có đánh đổi, chưa cải thiện toàn diện'],['RQ4: đóng góp riêng MSE','Ablation seed42 CI chứa0','Chưa đủ bằng chứng về lợi ích riêng']])
    para('SciFact là general retrieval, không kiểm tra trực tiếp temporal QA. Dữ liệu mined nhỏ; exact-text split chưa loại mọi trùng bài/đoạn trích. Chỉ ba seed train và ablation một seed. Chưa chạy temporal gold benchmark độc lập, nội suy time embedding, quy mô/thiết lập gốc hoặc đo hệ thống triệu tài liệu.')
    para('SciFact đã được xem trong thử nghiệm cũ để nhận diện mức giảm rồi mới thiết kế retention. Cố định kế hoạch lần chạy mới không loại bỏ thiên lệch thiết kế do đã biết benchmark. Cần đánh giá thêm benchmark chưa dùng khi thiết kế. Do các giới hạn này, nhóm chưa khẳng định thành công toàn diện, zero catastrophic forgetting hoặc chứng minh trọn vẹn bài báo; các quan sát kỹ thuật được xem là kết quả thử nghiệm, chưa phải phát minh đã được xác nhận.')
    heading('CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN',1)
    heading('5.1. Những phần nhóm đã thực hiện và giá trị kết quả')
    para('Nhóm đã chuyển từ demo giả lập và smoke test sang một quy trình dữ liệu thật có kiểm soát, đánh giá đầy đủ SciFact, đối chứng nhiều seed, lưu bằng chứng, bổ sung teacher/retention loss và kiểm tra checkpoint inference thực. Đây là phần bổ sung có thể kiểm tra lại so với báo cáo cũ; không chỉ là đề xuất chưa triển khai.')
    para('Kết luận chính: fixed-teacher retention cải thiện nDCG@10 so với TRPO đối chứng trên SciFact trong ba seed đã chạy, nhưng chưa vượt Contriever gốc và làm giảm mined-preference accuracy. MSE chưa có bằng chứng lợi ích riêng trong ablation hiện tại. Vì vậy có thể mô tả đây là biến thể hạn chế suy giảm retrieval tổng quát, chưa phải cải tiến toàn diện temporal retrieval hoặc đóng góp học thuật mới được peer review.')
    heading('5.2. Các phần chưa thực hiện và hướng tiếp theo')
    for text in [
      'Ưu tiên có temporal gold test độc lập và kiểm tra chất lượng mined labels; tránh tối ưu một proxy chưa đáng tin hoặc nhầm similarity với tính đúng mốc.',
      'Đánh giá benchmark chưa dùng khi thiết kế, mở rộng seed/ablation và thiết lập validation riêng để chọn hệ số; không dùng lại test để chọn mô hình tốt nhất rồi tuyên bố tổng quát.',
      'Nghiên cứu retention có điều kiện để tránh ép student giữ ưu tiên lỗi thời của teacher. Đây mới là hướng nghiên cứu, chưa có cài đặt/kết quả trong báo cáo này.',
      'Các ý tưởng LoRA/time adapters, event graphs, query-only temporal transform và isotropy/whitening vẫn là hướng mở; chưa có code/benchmark để khẳng định độ trễ, khả năng chống collapse hoặc lợi ích của chúng.',
      'Tái lập đầy đủ temporal QA, nhiều mốc và nội suy time embedding của bản bài báo được chọn khi có dữ liệu/tài nguyên; đo chi phí lập chỉ mục theo đúng hệ thống triển khai.'
    ]:para(text,'List Bullet')
    heading('5.3. Đối chiếu mục tiêu với kết quả đạt được')
    table(['Mục tiêu','Mức đạt và căn cứ','Điều chưa thể kết luận'],[
      ['Hiểu/cài đặt phương pháp','Có bản CE/TRPO, EMA, queue và 14 kiểm tra toàn bộ pipeline','Không đồng nghĩa đã tái lập nguyên bài báo'],
      ['Thay dữ liệu mock','Có 2.992 bộ ba thật; train/held-out riêng và provenance','Nhãn mined chưa thành temporal gold; split chưa loại mọi trùng ngữ nghĩa'],
      ['Thử một mở rộng có code','Có teacher cố định, hai loss giữ kiến thức, ba seed và ablation','Không phải phát minh distillation; chưa chứng minh mọi thành phần cần thiết'],
      ['Cải thiện retrieval sau fine-tune','nDCG@10 tăng so với TRPO ở cả ba seed đã chạy','Chưa vượt model gốc, chưa chứng minh tổng quát sang benchmark khác'],
      ['Minh họa sử dụng thật','Có inference checkpoint, query mới, cache và đối chiếu top-k','Một query demo không chứng minh accuracy hoặc temporal correctness']])
    para('Đóng góp có giá trị nhất của đồ án là biến một minh họa thuận lợi thành thực nghiệm có đối chứng và có thể truy nguyên. Phần kỹ thuật và phần nghiên cứu cần được đánh giá riêng: pipeline hoạt động là điều kiện cần để tin số liệu; hiệu quả phương pháp phải dựa vào benchmark, độ bất định và các kết quả bất lợi. Nhóm đạt được bằng chứng cho một giả thuyết hẹp về retention, chưa đạt mục tiêu cải thiện toàn diện TPOUR.')
    heading('5.4. Ý nghĩa kết quả và những cách giải thích còn cạnh tranh')
    para('SciFact tăng cùng top-10 overlap với baseline phù hợp với cơ chế giữ hành vi teacher. Tuy nhiên, tương quan này chưa chứng minh nguyên nhân duy nhất; regularization, thay đổi quỹ đạo tối ưu hoặc nhãn train có nhiễu cũng có thể ảnh hưởng. Kết quả gợi ý cần giữ đối chứng chỉ KL, đồng thời mở rộng seed và kiểm tra benchmark khác trước khi chọn thiết kế phức tạp hơn.')
    para('Preference accuracy giảm là bằng chứng biến thể đang đánh đổi một phần mục tiêu mined để giữ retrieval tổng quát. Do nhãn chưa được thẩm định, chưa biết mức giảm phản ánh mất khả năng đúng thời gian hay chỉ không khớp proxy. Cả hai khả năng đều cần được xem xét: không thể bỏ qua proxy xấu, cũng không thể gọi proxy là chân lý về thời gian.')
    para('Việc chưa vượt Contriever gốc cho thấy fine-tune hiện tại chưa tạo lợi ích ròng trên SciFact. Với tác vụ chỉ truy xuất khoa học tổng quát, chưa có căn cứ chọn retention thay cho baseline từ số liệu này. Với tác vụ cần thời gian, cần thêm tập temporal gold để đánh giá lợi ích cập nhật so với chi phí giữ kiến thức; SciFact một mình không đủ để quyết định.')
    heading('5.5. Các rào cản còn lại và tác động đến độ tin cậy')
    table(['Rào cản','Tác động','Cách xử lý cần làm tiếp'],[
      ['Mẫu train nhỏ và mined labels','Khó khái quát; accuracy thấp có nhiều cách giải thích','Kiểm tra thủ công nhãn và mở rộng dữ liệu có nguồn/thời gian rõ'],
      ['Thiếu article ID và temporal gold','Exact-text split chưa loại mọi leakage; chưa đo đúng mốc','Nhóm theo article/snapshot khi có metadata; xây test thời gian được thẩm định'],
      ['Đã dùng SciFact khi thiết kế','Có benchmark-aware design bias dù không chọn checkpoint bằng test','Giữ benchmark mới chưa xem để đánh giá xác nhận'],
      ['Ba seed, ablation một seed','CI theo query không bao quát mọi nguồn bất định','Lặp nhiều seed cho mọi nhánh, báo CI/SD đúng đơn vị'],
      ['1.000 lần thử và cấu hình gọn','Chưa đại diện quá trình hội tụ hay cấu hình bài gốc','Theo dõi validation độc lập và so quy mô/lịch train có kiểm soát'],
      ['Chi phí và tái lập môi trường','Teacher tăng train/memory; checkpoint lớn nằm ở Kaggle','Lưu phiên bản/hash và hướng dẫn tải; đo chi phí end-to-end khi mở rộng']])
    para('Những rào cản này giới hạn phạm vi tuyên bố, không làm mất giá trị của toàn bộ thực nghiệm. Nhóm vẫn có thể chứng minh code chạy, số liệu có nguồn và biến thể có tác động đo được trong cấu hình cụ thể. Điều cần tránh là dùng kết quả đó để khẳng định mọi temporal retrieval đều được cải thiện hoặc mô hình không còn quên tri thức.')
    heading('5.6. Kế hoạch phát triển theo thứ tự ưu tiên và tiêu chí kiểm chứng')
    table(['Ưu tiên','Thí nghiệm đề xuất — chưa thực hiện','Tiêu chí quyết định'],[
      ['1. Kiểm tra mục tiêu thời gian','Thẩm định mẫu mined với người đọc, ghi nhãn đúng/sai/không đủ thông tin; tách train/validation/test theo nguồn và thời gian','Có test temporal gold độc lập, quy tắc nhãn rõ và thống kê bất đồng'],
      ['2. Xác nhận tính tổng quát','Cố định cấu hình trên validation; đánh giá benchmark chưa dùng khi thiết kế; lặp ablation KL và KL+MSE nhiều seed','Lợi ích lặp lại và báo cả kết quả giảm; không chọn tham số trên test'],
      ['3. Giảm đánh đổi retention','Thử trọng số retention có điều kiện khi teacher mâu thuẫn với ưu tiên mới; so gamma=0, chỉ KL và biến thể điều kiện','Đánh giá đồng thời temporal quality, retrieval tổng quát và chi phí; loại nếu chỉ chuyển suy giảm sang mục tiêu khác'],
      ['4. Mở rộng gần bài gốc','Dữ liệu nhiều snapshot, kiểm tra nội suy time embedding và lịch train dài hơn','Đối chiếu đúng bản PDF/protocol, báo độ lệch cấu hình và tài nguyên']])
    para('Retention có điều kiện là một giả thuyết nghiên cứu tiếp theo, chưa có kết quả trong kho hiện tại. Không nên gán trọng số bằng nhãn test hoặc quyết định theo trực giác về một query; cần định nghĩa tín hiệu tại train/validation và giữ test độc lập. Chỉ khi có lợi ích đồng thời trên hai mục tiêu và chi phí chấp nhận được mới có cơ sở nói phương pháp toàn diện hơn.')
    heading('5.7. Các câu hỏi trọng tâm khi bảo vệ đồ án')
    for question,answer in [
      ('Nhóm cải tiến gì so với việc chạy lại code?', 'Nhóm thêm teacher Contriever cố định và objective giữ phân bố xếp hạng/biểu diễn, rồi kiểm chứng bằng đối chứng ba seed, ablation và checkpoint thật. Runner/provenance là đóng góp kỹ thuật hỗ trợ kiểm chứng; không nên gọi riêng việc chuyển lên Kaggle là cải tiến thuật toán.'),
      ('Nếu chưa vượt baseline thì cải tiến có ý nghĩa gì?', 'Ý nghĩa hiện có là hạn chế mức giảm do fine-tune so với TRPO trong cùng protocol. Đây là kết quả có giới hạn; nếu chỉ cần SciFact thì baseline vẫn có điểm cao hơn. Chưa có bằng chứng để khuyên thay baseline trong mọi ứng dụng.'),
      ('Vì sao không dùng EMA reference làm teacher luôn?', 'EMA reference thay đổi theo student để phục vụ preference learning. Teacher cố định giữ mốc hành vi trước fine-tune. Nếu dùng chung, mục tiêu retention sẽ di chuyển theo quá trình học và không còn đúng thiết kế đã thử.'),
      ('Có chắc MSE cần thiết và mô hình tốt hơn về thời gian?', 'Chưa chắc cả hai. Ablation một seed có CI chứa 0; preference proxy giảm, chưa có temporal gold. Nhóm báo rõ các kết quả này thay vì kết luận mọi thành phần đều hiệu quả.'),
      ('Có tái lập đầy đủ bài báo chưa?', 'Chưa. Nhóm đã cài đặt bản gọn, train dữ liệu thật, đo SciFact và thử retention; chưa tái lập toàn bộ snapshot, temporal QA, time interpolation và quy mô tác giả.'),
      ('Demo chứng minh điều gì?', 'Demo xác nhận checkpoint thực có thể mã hóa query và truy xuất corpus, với hash/top-k đối chiếu. Chất lượng được đánh giá trên toàn test; query mới chưa có gold nên chỉ minh họa hoạt động, không phải bằng chứng accuracy.')
    ]:
        x=para(question);x.runs[0].bold=True;para(answer)
    heading('PHỤ LỤC: MÃ NGUỒN, ĐẦU RA VÀ TÀI LIỆU THAM KHẢO',1)
    heading('A. Nguồn bằng chứng và lệnh kiểm tra')
    para('Kho nhóm: https://github.com/lequangduccode/tpour-trpo-experiments. Source train retention: 219a21a; source live demo: 0d2e084. Commit source cố định được lưu trong executed_source.ipynb/provenance, nên các thay đổi hình thức báo cáo sau đó không được nhầm với code dùng train.')
    table(['Đường dẫn','Nội dung'],[
      ['results/kaggle_real_seed42/','Giai đoạn real-data CE/TRPO một seed; khác protocol lịch sử 30 query'],['results/anchored_kaggle/','Kết quả ba seed, ablation, config/log/metrics/per-query/top100'],['results/anchored_kaggle/verification.json','Tính lại metric/CI, đối chiếu config/hash và live evidence'],['results/anchored_kaggle/live_demo/','JSON inference thật, truy vấn mới, hash, log và notebook đã chạy'],['results/anchored_kaggle/demo/','Biểu đồ và ảnh demo tạo bằng code từ số liệu thật'],['Kaggle Output của run train','encoder.pt từng model; không có checkpoint lớn trong GitHub']])
    for cmd in [
      'python -m unittest test_tpour test_real_experiment test_anchored_tpour -v',
      'python verify_anchor_results.py --root results/anchored_kaggle',
      'python demo_retention.py --root results/anchored_kaggle --query-id 1',
      'python demo_retention.py --root results/anchored_kaggle --query-id 1282',
      'python demo_retention.py --root results/anchored_kaggle --query-id 314',
      'python live_retrieval_demo.py --checkpoint <encoder.pt tải từ Kaggle> --query-id 1 --device cuda',
      'python live_retrieval_demo.py --checkpoint <encoder.pt tải từ Kaggle> --query "What influences the mechanical properties of biological materials?" --device cuda']:
        x=para(cmd);x.paragraph_format.keep_together=True
        for rr in x.runs:rr.font.name='Consolas';rr.font.size=Pt(9)
    para('Lần đầu live demo mã hóa corpus và lưu cache; query sau có thể dùng cache nếu hash/protocol khớp. Cài requirements-experiment.txt; cần tải SciFact và checkpoint đúng run. Kaggle hiện dùng notebook thuộc tài khoản chủ repo thử nghiệm, không mặc nhiên cấp quyền cho người xem GitHub.')
    para('Notebook train: https://www.kaggle.com/code/vananh1911/tpour-retention-improvement-demo. Notebook live: https://www.kaggle.com/code/vananh1911/tpour-live-checkpoint-verification. Nếu chưa có quyền xem notebook riêng, đối chiếu JSON/log/ảnh và mã notebook được lưu công khai trong repo hoặc chạy lại trên tài khoản của mình.')
    heading('B. Cách đọc số liệu cũ')
    para('Số 0,74472 (TPOUR) và0,74413 (baseline) thuộc subset30 queries trong bản Hạnh. Giá trị Recall@10=0,90 không đại diện kết quả full300. Toy CE/TRPO cùng nDCG=1,0 chỉ là smoke test thuận lợi. Không so raw similarity giữa các model như một độ chính xác; scale embedding có thể khác nhau. Các dữ kiện này chỉ giữ như lịch sử kiểm chứng, không dùng làm kết luận chính của bản cập nhật.')
    heading('C. Tài liệu tham khảo')
    for text in [
      '[1] HyunJin Kim, Jaejun Shim, Young Jin Kim, JinYeong Bak. Temporal Preference Optimization for Unsupervised Retrieval. Bản PDF nhóm chọn ghi ICML2026/PMLR306, 33 trang; upstream/TPOUR/ICML_2026_TPOUR.pdf. Source: https://github.com/agwaBom/TPOUR.',
      '[2] Gautier Izacard et al. Towards Unsupervised Dense Information Retrieval with Contrastive Learning. https://arxiv.org/abs/2112.09118.',
      '[3] Nandan Thakur et al. BEIR: A Heterogeneous Benchmark for Zero-shot Evaluation of Information Retrieval Models. https://arxiv.org/abs/2104.08663; https://github.com/beir-cellar/beir.',
      '[4] Rafael Rafailov et al. Direct Preference Optimization: Your Language Model is Secretly a Reward Model. https://arxiv.org/abs/2305.18290.',
      '[5] Kaiming He et al. Momentum Contrast for Unsupervised Visual Representation Learning. CVPR2020.',
      '[6] Zhizhong Li, Derek Hoiem. Learning without Forgetting. https://arxiv.org/abs/1606.09282.',
      '[7] Weighted KL-Divergence for Document Ranking Model Refinement. https://arxiv.org/abs/2406.05977.',
      '[8] SciFact dataset và qrels: https://github.com/allenai/scifact; bản BEIR được lưu hash trong scifact_download.json.'
    ]:para(text)
    # Remove orphaned old pictures from main-document relationships, keep active code-generated figures.
    used={e.get(qn('r:embed')) for e in doc._element.iter() if e.get(qn('r:embed'))}
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
    # Normalize spacing in prose without changing code paths/commands.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for old,new in [('khoảng0','khoảng 0'),('khoảng332','khoảng 332'),('và385','và 385'),('dưới0','dưới 0'),('chứa0','chứa 0'),('bỏ5','bỏ 5'),('bằng0','bằng 0'),('subset30','subset 30'),('full300','full 300'),('full300','full 300'),('thang100','thang 100')]:
                if old in run.text:run.text=run.text.replace(old,new)
    # Font sanity for all editable math, including preserved template equations.
    for math in doc._element.xpath('.//m:oMath'):
        for run in math.iter(qn('m:r')):
            props=run.find(qn('w:rPr'))
            if props is None:props=OxmlElement('w:rPr');run.insert(0,props)
            fonts=props.find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');props.append(fonts)
            for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Cambria Math')
    # Replace inherited manual page breaks with deliberate section boundaries.
    for paragraph in doc.paragraphs:
        for br in list(paragraph._p.xpath('.//w:br')):
            if br.get(qn('w:type'))=='page':br.getparent().remove(br)
        if paragraph.style.name=='Heading 1':paragraph.paragraph_format.page_break_before=True
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.paragraph_format.keep_with_next=True
    p[27].paragraph_format.page_break_before=True
    for t in doc.tables:
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if i==0 and pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
            if len(t.rows)<=6 and i<len(t.rows)-1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
    doc.core_properties.title='Báo cáo TPOUR/TRPO cập nhật Kaggle và retention'
    doc.core_properties.subject='Thực nghiệm dữ liệu thật, biến thể objective và các đánh đổi quan sát được'
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    assert sha(source)==before,'Original document changed unexpectedly'
    manifest={'original':str(source),'original_sha256':before,'output':str(output),'output_sha256':sha(output),
              'format':'Original cover, assignment placeholders, chapter1/2 styles and native equations retained; chapters3–5 and appendices rewritten from verified evidence.',
              'source_train_commit':'219a21a','source_live_commit':'0d2e084','original_unchanged':True}
    (output.parent/'report_update_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))


if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--source',required=True);a.add_argument('--output',required=True);v=a.parse_args();build(v.source,v.output)
