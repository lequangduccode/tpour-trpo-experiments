"""Rebuild on the exact PTIT template; retain native math and import verified new evidence."""
import copy
import hashlib
import io
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from update_main_report_after_kaggle import mn, mr, seq, sub, sup, frac

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'reports/ptit_updated/Bao_cao_TPOUR_TRPO_da_sua_cong_thuc.docx'
UPDATE=ROOT/'reports/revised_main/Bao_cao_TPOUR_TRPO_cap_nhat_Kaggle_retention.docx'
OUT=ROOT/'reports/ptit_verified'
NAME='Bao_cao_TPOUR_TRPO_PTIT_cap_nhat_da_ra_soat'
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def text(p):return ''.join(t.text or '' for t in p.iter(qn('w:t')))
def summ(lower,upper,body):
    pr=mn('naryPr',mn('chr'),mn('limLoc'))
    pr[0].set(qn('m:val'),'∑');pr[1].set(qn('m:val'),'undOvr')
    if upper is None:
        e=mn('supHide');e.set(qn('m:val'),'1');pr.append(e)
    return mn('nary',pr,mn('sub',seq(lower)),mn('sup',seq(upper or '')),mn('e',seq(body)))
def clear(p):
    for c in list(p._p):
        if c.tag!=qn('w:pPr'):p._p.remove(c)
def replace(p,t):clear(p);p.add_run(t)
def equation(p,parts):
    clear(p);prop=mn('oMathParaPr');jc=mn('jc');jc.set(qn('m:val'),'center');prop.append(jc)
    p._p.append(mn('oMathPara',prop,mn('oMath',parts)))
    p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1.15
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
def append_eq(doc,parts,caption):
    equation(doc.add_paragraph(),parts);doc.add_paragraph(caption,'Caption')


def build():
    before=digest(SOURCE);doc=Document(SOURCE);old=list(doc.paragraphs)
    assert len(old)==241 and old[90].text.startswith('CHƯƠNG 3:')
    # Keep the actual cover, roster, section layout, chapter 1 and chapter 2.
    replacing={
      23:'Báo cáo dựa trên bài báo “Temporal Preference Optimization for Unsupervised Retrieval” của HyunJin Kim, Jaejun Shim, Young Jin Kim và JinYeong Bak, bản PDF nhóm đã chọn [1]. Bài báo nghiên cứu temporal misalignment, kết hợp contrastive learning với Temporal Retrieval Preference Optimization (TRPO), dùng reference momentum/queue và nội suy time vector. Nội dung kế thừa từ tác giả được phân biệt với biến thể retention do nhóm bổ sung.',
      24:'Bản thực hành đã chuyển từ demo toy và log 20 bước Contriever của giai đoạn trước sang huấn luyện trên mẫu Wikipedia thật bằng Kaggle T4. Nhóm làm sạch 3.000 dòng còn 2.992 bộ ba, tách 2.543 train/449 held-out preference và đánh giá đủ 300 truy vấn SciFact trên 5.183 tài liệu. Code bổ sung fixed teacher, ranking KL và relative MSE nhằm kiểm tra khả năng hạn chế suy giảm retrieval sau fine-tune.',
      25:'Qua ba seed 42/43/44, CE+TRPO đạt nDCG@10 trung bình 0,60035 ± 0,00432; biến thể retention đạt 0,62907 ± 0,00204; Contriever gốc đạt 0,63935. Retention tốt hơn TRPO nhưng chưa vượt model gốc; mined-preference accuracy giảm từ 0,44915 xuống 0,42910. Vì vậy kết luận chỉ là hạn chế suy giảm truy xuất tổng quát trong cấu hình đã đo, chưa phải cải thiện toàn diện truy xuất theo thời gian.',
      30:'Phần thực hành có kiểm tra code, provenance, rankings/metrics, ba seed đối chứng, ablation một seed và demo inference từ checkpoint. SciFact là benchmark retrieval tổng quát; held-out preference dùng nhãn mined chưa được thẩm định thành temporal gold. Kết quả chưa tái lập temporal QA, time interpolation nhiều mốc hoặc timestamp predictor của bài báo. Những điểm này được đối chiếu ở Chương 3–5.',
      34:'Truy xuất thông tin xếp hạng tài liệu theo độ liên quan với truy vấn. BM25 dựa nhiều vào từ vựng; dense retrieval học biểu diễn ngữ nghĩa. Tuy nhiên, kho nhiều năm có thể chứa các phiên bản cùng chủ đề nhưng khác nội dung. Khi hỏi quy trình hiện tại, văn bản cũ vẫn có thể gần ngữ nghĩa nhưng sai mốc cần tìm. TPOUR nghiên cứu sự lệch giữa thời gian truy vấn và tài liệu này, không chỉ việc tìm văn bản có nhiều từ giống nhau [1].',
      38:'Cài đặt và kiểm tra CE/TRPO, reference momentum, queue; phân biệt tính đúng của code với chất lượng retrieval.',
      39:'Giữ đối chứng CE-only của giai đoạn đầu; ở giai đoạn retention mới, so CE+TRPO và biến thể cùng cấu hình, ba seed, kèm ablation chỉ KL ở seed42.',
      40:'Thực nghiệm dữ liệu tác giả trên Kaggle, đánh giá đủ SciFact, lưu cấu hình/log/hash và kiểm chứng inference bằng checkpoint thật.',
      41:'Phân tích mức đạt mục tiêu, đánh đổi, chi phí và khoảng cách tới thiết lập temporal retrieval của tác giả.',
      42:'Toy chỉ kiểm tra pipeline. Thí nghiệm mới train Wikipedia và đánh giá SciFact với vai trò tách biệt; chưa có temporal gold để kết luận mô hình hiểu đúng mọi mốc thời gian.',
      48:'Mục 4 của bài báo đánh giá SituatedQA, RealTimeQA và BEIR [1, mục 4.1–4.4]. Nhóm đã có full SciFact với ba seed cho đối chứng retention, nhưng chưa tái lập hai benchmark temporal QA, timestamp predictor hoặc tổng quát hóa mốc trung gian. Kết quả của tác giả và kết quả của nhóm vì thế được trình bày riêng.',
      49:'TRPO và CE/MoCo được triển khai gọn trong tpour.py; phép nội suy tensor có trong lệnh interpolate. Pipeline dữ liệu thật/đánh giá ở real_experiment.py và teacher/retention ở anchored_tpour.py. Phần nội suy chưa được kiểm chứng trên các checkpoint chuyên biệt thời gian thật; mixture-of-TPOUR chưa triển khai. Upstream được giữ để đối chiếu các khác biệt kỹ thuật.',
      68:'B là batch size; i chỉ bộ ba trong batch; σ là sigmoid. D⁺ và D⁻ là preferred/rejected theo nhãn mined, chưa mặc nhiên là temporal gold. Sθ(Q,D) trong mục này là dot product chia nhiệt độ T; điểm retrieval s(Q,D) ở mục 2.2 là dot product chưa chia T.',
      69:'Phương trình (4) của bài báo so sánh chênh lệch preferred/rejected giữa current và reference [1, mục 3.1]. Code nhóm và upstream đã lưu đều chia nhiệt độ trong phần score TRPO; code nhóm dùng T=0,05. Bỏ reference sẽ đổi objective. Khi chênh lệch tương đối tăng, log-sigmoid loss giảm; kiểm tra gradient chỉ xác nhận những trường hợp tensor đã thử.',
      72:'𝒬 là tập chỉ số negative keys trong queue; kⱼ là vector negative reference đã lưu. Trong biểu thức CE, mỗi tích vô hướng đều được chia cùng nhiệt độ T. Code lấy trung bình trên batch; qᵢ và pᵢ lần lượt là embedding current của query/preferred. Không dùng qrels SciFact để tạo positive khi train.',
      74:'Phương trình (3) của bài báo mô tả CE dùng negative embeddings từ queue; mục 3.1 kết hợp CE/TRPO và cập nhật reference bằng momentum [1, tr. 4; Hình 2]. Bản thực hành dùng λ=0,925, β=0,3 và T=0,05. Đây là trọng số loss, không phải tỷ lệ tài liệu đúng mốc. Bản gọn và upstream còn khác chi tiết CE/in-batch negatives, chế độ reference và lịch/thứ tự tính; không coi đây là tái lập bit-for-bit.',
      77:'Các khác biệt cài đặt CE, reference, queue, scheduler và checkpoint được tổng hợp tại mục 3.11. Bảng phân biệt bản gọn hiện tại, upstream và những nhận xét chỉ còn đúng với code cũ.',
    }
    for i,t in replacing.items():replace(old[i],t)
    # Make the CE denominator explicit and remove ambiguous negative score notation.
    dotpos=seq(sup(sub('q','i'),'⊤'),sub('p','i'))
    dotneg=seq(sup(sub('q','i'),'⊤'),sub('k','j'))
    exppos=seq('exp(',frac(dotpos,'T'),')')
    expneg=seq('exp(',frac(dotneg,'T'),')')
    equation(old[71],seq(sub('ℒ','CE'),' = −',frac('1','B'),summ('i = 1','B',seq('ln ',frac(copy.deepcopy(exppos),seq(copy.deepcopy(exppos),' + ',summ('j ∈ 𝒬',None,expneg)))))))
    # Temperature consistency and preserved interpolation are checked later structurally.
    start=old[90]._p;removing=False
    for e in list(doc._element.body):
        if e is start:removing=True
        if removing and e.tag!=qn('w:sectPr'):doc._element.body.remove(e)
    # The first two roster tables are inherited; no invented reassignment of authors.
    evidence=doc.tables[2]
    for row in list(evidence.rows)[1:]:evidence._tbl.remove(row._tr)
    for row in [
      ['Toy + lịch sử Hạnh','Toy 3 seed; Contriever 20 bước và SciFact 30/300 query','Bằng chứng lịch sử; không dùng thay full test'],
      ['Train dữ liệu thật mới','2.543 bộ ba train; ba seed CE+TRPO/retention; 995 cập nhật thực mỗi run','Có đối chứng cùng protocol trên Kaggle'],
      ['Retrieval / proxy','Full SciFact; 449 bộ ba mined giữ riêng','General retrieval tăng so với TRPO, proxy giảm'],
      ['Kiểm tra / demo','14 tests; kiểm tra metrics/hash; inference checkpoint thật','Xác nhận các kiểm tra đã thực hiện, không phải temporal gold']]:
        for c,v in zip(evidence.add_row().cells,row):c.text=v
    terms=doc.tables[3]
    for term,meaning in [('Fixed teacher','Contriever khởi tạo đóng băng; khác reference EMA'),('Retention','Hạng loss giữ phân bố xếp hạng và biểu diễn của teacher'),('Ablation','Thử bỏ thành phần để kiểm tra tác dụng riêng'),('Mined preference','Nhãn ưu tiên khai thác từ dữ liệu; chưa phải temporal gold')]:
        for c,v in zip(terms.add_row().cells,[term,meaning]):c.text=v
    # Keep PTIT styling; add missing named styles for newly imported paragraphs.
    if 'Caption' not in doc.styles:
        st=doc.styles.add_style('Caption',WD_STYLE_TYPE.PARAGRAPH);st.font.name='Times New Roman';st.font.size=Pt(11);st.font.italic=True
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.name='Times New Roman'
        doc.styles[name].font.size=Pt(16 if name=='Heading 1' else 13)
        doc.styles[name].font.bold=True
    doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(12)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.add_heading('2.7. Độ đo và quy tắc tổng hợp kết quả',2)
    doc.add_paragraph('SciFact dùng qrels nhị phân. Các độ đo dưới đây được tính cho từng truy vấn rồi lấy trung bình trên đủ 300 truy vấn có nhãn; truy vấn không có hit được giữ với điểm 0. Một nDCG@10 bằng 0,63 không phải 63% câu trả lời đúng. Loss train, raw similarity và retrieval metrics là các đại lượng khác nhau.')
    for i in [155,156,157,158,159,160]:doc._element.body.insert(len(doc._element.body)-1,copy.deepcopy(old[i]._p))
    append_eq(doc,seq('AP@k(q) = ',frac('1','|𝓡q|'),summ('i = 1','k',seq('𝟙[',sub('d','q,i'),' ∈ 𝓡q] · P@i(q)'))),'AP chia tổng số tài liệu liên quan của query, kể cả khi k nhỏ hơn số positive.')
    append_eq(doc,seq('MAP@k = ',frac('1','|𝒰|'),summ('q ∈ 𝒰',None,'AP@k(q)')),'𝒰 là tập query; precision P@i(q) là số positive trong top-i chia i.')
    doc.add_paragraph('IDCG dùng thứ tự lý tưởng trên cùng số positive và cutoff. Với qrels nhị phân, gain 2^rel−1 bằng 1 ở mỗi positive; công thức DCG tổng quát vì thế trùng cách tính binary_metrics trong real_experiment.py. MRR lấy reciprocal rank của hit đầu tiên trong top-k; Recall có thể khác hit-rate nếu query có nhiều positive. SD qua ba seed và CI bootstrap theo query đo hai nguồn bất định khác nhau.')
    doc.add_heading('2.8. Cơ sở cho phần mở rộng retention của nhóm',2)
    doc.add_paragraph('Learning without Forgetting [7] là nền tảng về giữ năng lực cũ; KL ranking distillation là hướng liên quan [8]. TPOUR [1] vẫn là nguồn chính. Nhóm dùng KL tiêu chuẩn và relative MSE trong anchored_tpour.py, không cài weighted-KL của [8].')
    doc.add_paragraph('Teacher cố định nhằm giữ retrieval nhưng có thể cản học ưu tiên mới; cần đo cả SciFact và preference proxy. Chi tiết thiết kế ở Chương 3.')
    # Import only the implementation/results/discussion, with all image relations remapped.
    newer=Document(UPDATE);begin=False;qa=[];in_qa=False;refs=False
    for e in newer._element.body:
        t=text(e)
        if t.startswith('CHƯƠNG 3:'):begin=True
        if not begin or e.tag==qn('w:sectPr'):continue
        if t.startswith('5.7. Các câu hỏi'):in_qa=True;continue
        if in_qa and t.startswith('PHỤ LỤC:'):in_qa=False
        if in_qa:qa.append(copy.deepcopy(e));continue
        if t.startswith('C. Tài liệu tham khảo'):refs=True;break
        if t.startswith('Các ý tưởng LoRA/time adapters'):continue
        clone=copy.deepcopy(e)
        # Citations use the PTIT source bibliography numbering.
        for node in clone.iter(qn('w:t')):
            if node.text:node.text=node.text.replace('BEIR [3]','BEIR [4]')
        for node in clone.iter():
            for attr in [qn('r:embed'),qn('r:link'),qn('r:id')]:
                rid=node.get(attr)
                if rid and rid in newer.part.rels:
                    rel=newer.part.rels[rid]
                    if rel.reltype.endswith('/image'):
                        nrid,_=doc.part.get_or_add_image(io.BytesIO(rel.target_part.blob))
                    elif rel.is_external:nrid=doc.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
                    else:raise ValueError('Unexpected internal relation while importing '+rel.reltype)
                    node.set(attr,nrid)
        doc._element.body.insert(len(doc._element.body)-1,clone)
    # Add a real code diagram near the encoder-role explanation (Figure 1).
    marker=next(p for p in doc.paragraphs if p.text.startswith('3.3. Ba vai trò'))
    pic=doc.add_paragraph();pic.add_run().add_picture(str(ROOT/'reports/anchored_improvement/method.png'),width=Inches(6.1))
    cap=doc.add_paragraph('Hình 1. Bản gọn TPOUR và fixed teacher bổ sung; sơ đồ vẽ bằng code.','Caption')
    marker._p.addnext(pic._p);pic._p.addnext(cap._p)
    # Explicit normalized candidate probabilities, not an unexplained KL symbol.
    marker=next(p for p in doc.paragraphs if p.text.startswith('Relative representation loss'))
    new=doc.add_paragraph();equation(new,seq(sub('P','a,i,j'),' = ',frac(seq('exp(cos(',sub('q','a,i'),', ',sub('d','a,j'),') / ',sub('T','r'),')'),seq(summ('ℓ = 1','2B',seq('exp(cos(',sub('q','a,i'),', ',sub('d','a,ℓ'),') / ',sub('T','r'),')'))))))
    marker._p.addprevious(new._p)
    caption=doc.add_paragraph('Với a là student hoặc teacher, j∈{1,…,2B}, Tr=0,1; candidates gồm B preferred và B rejected.','Caption');marker._p.addprevious(caption._p)
    new=doc.add_paragraph();equation(new,seq(sub('ℒ','KL'),' = ',frac('1','B'),summ('i = 1','B',summ('j = 1','2B',seq(sub('P','teacher,i,j'),' ln ',frac(sub('P','teacher,i,j'),sub('P','student,i,j')))))))
    marker._p.addprevious(new._p)
    # Restore the technical comparison from the old report, corrected for today's runner.
    marker=next(p for p in doc.paragraphs if p.text.startswith('CHƯƠNG 4:'))
    blocks=[]
    blocks.append(doc.add_heading('3.11. Khác biệt kỹ thuật với upstream và bản thực hành cũ',2)._p)
    blocks.append(doc.add_paragraph('So sánh này bám upstream/TPOUR/src/moco.py tại commit d1d8e70 và code thực hành dùng trong các run đã lưu. Nó giải thích vì sao cùng tên TPOUR không đồng nghĩa cùng mọi chi tiết cài đặt. Các giá trị quy mô và số liệu bài báo vẫn phải đối chiếu đúng cấu hình tác giả; nhóm không dùng bảng dưới để quy kết nguyên nhân duy nhất của chênh lệch chất lượng.')._p)
    comparison=doc.add_table(rows=1,cols=3)
    for c,v in zip(comparison.rows[0].cells,['Chi tiết','Upstream đã lưu','Bản gọn / runner hiện tại']):c.text=v
    for row in [
      ['CE candidates / positive','Ma trận query-current với preferred-reference toàn batch, queue và weak positives khác đường chéo','Positive query/preferred đều current; chỉ queue làm negatives CE'],
      ['Nhiệt độ và smoothing','Score TRPO chia T; CE scores ở forward này không chia T, có tùy chọn label smoothing','CE và TRPO cùng chia T=0,05; không label smoothing'],
      ['Reference / thời điểm EMA','Cập nhật EMA trong forward train trước khi mã hóa reference; có tùy chọn train mode','Reference eval; EMA sau optimizer step hợp lệ'],
      ['Queue','Gather đa tiến trình, yêu cầu queue chia hết số keys được enqueue; khởi tạo ngẫu nhiên trong lớp','Buffer vòng xử lý wrap bất kỳ; warmup bằng train embeddings; một GPU mỗi run'],
      ['Scheduler / checkpoint','Hệ train đầy đủ thuộc code tác giả','Runner mới có warmup100/linear decay và trạng thái resume; nhận xét cũ “LR cố định/chỉ lưu encoder” không còn áp dụng'],
      ['KL / teacher','Có tính kl_loss giữa q_q và k_p để ghi log; không cộng vào loss đang trả về','Teacher cố định riêng; ranking KL trên candidates + relative MSE được cộng vào objective retention']]:
        for c,v in zip(comparison.add_row().cells,row):c.text=v
    blocks.append(comparison._tbl)
    for e in blocks:marker._p.addprevious(e)
    # Keep report discussion academic; defense prompts belong in an appendix.
    doc.add_heading('C. Gợi ý giải thích khi bảo vệ',2)
    for e in qa:doc._element.body.insert(len(doc._element.body)-1,e)
    doc.add_heading('D. Ranh giới dữ liệu và kiểm tra nguồn context',2)
    doc.add_paragraph('Trong pipeline đã chạy, “context” là trường văn bản trong JSONL của tác giả, không phải nội dung lấy thêm từ một chatbot. clean_rows chỉ lấy question, context đầu tiên của positive_ctxs và weak_positive_ctxs; chuẩn hóa phục vụ lọc/chia tập rồi train chỉ đọc train.jsonl. Queue chứa embeddings từ chính train split; teacher là bản sao model khởi tạo, không truy cập một kho văn bản ngoài. Word deletion chỉ xóa từ, không sinh thêm thông tin.')
    doc.add_paragraph('SciFact là dữ liệu ngoài mẫu Wikipedia nhưng được dùng có chủ đích làm corpus/query/qrels đánh giá tổng quát, không làm bộ ba cập nhật gradient. Demo truy xuất title+text trong đúng corpus SciFact; nó không gọi web search hoặc LLM để thêm context hay sinh câu trả lời. Contriever/tokenizer tải từ mô hình tiền huấn luyện là phụ thuộc có khai báo; nhóm không có khả năng kiểm toán toàn bộ dữ liệu tiền huấn luyện của mô hình này.')
    doc.add_paragraph('Không thể từ cấu trúc pipeline suy ra không có mọi dạng nhiễm benchmark. SciFact đã được xem khi thiết kế retention, và split chỉ kiểm tra trùng exact text trong mẫu train/held-out. Báo cáo không khẳng định đã loại trùng ngữ nghĩa hoặc bảo đảm dữ liệu tiền huấn luyện không từng chứa nội dung liên quan. File audit kèm báo cáo ghi các hash, kiểm tra và giới hạn cụ thể.')
    doc.add_heading('TÀI LIỆU THAM KHẢO',1)
    refs=[
      '[1] HyunJin Kim, Jaejun Shim, Young Jin Kim, JinYeong Bak. Temporal Preference Optimization for Unsupervised Retrieval. Bản PDF nhóm chọn, 33 trang, ghi ICML 2026/PMLR 306; upstream/TPOUR/ICML_2026_TPOUR.pdf. Google Drive: https://drive.google.com/file/d/1NNPeLJCRIj22_sZoAyJfDye1PzEpT9lz/view. Trích mục 3.1–3.3, phương trình (3)–(5), Hình 2 và mục 4.',
      '[2] agwaBom/TPOUR. Mã tác giả tại commit d1d8e70326deae6ef2ea03a8cb2c0cbc60db7a78, lưu trực tiếp trong upstream/TPOUR. https://github.com/agwaBom/TPOUR.',
      '[3] Gautier Izacard et al. Towards Unsupervised Dense Information Retrieval with Contrastive Learning. https://arxiv.org/abs/2112.09118. Model: https://huggingface.co/facebook/contriever.',
      '[4] Nandan Thakur et al. BEIR: A Heterogeneous Benchmark for Zero-shot Evaluation of Information Retrieval Models. https://arxiv.org/abs/2104.08663. https://github.com/beir-cellar/beir.',
      '[5] Rafael Rafailov et al. Direct Preference Optimization: Your Language Model is Secretly a Reward Model. https://arxiv.org/abs/2305.18290.',
      '[6] Bản code Hạnh dùng ở giai đoạn lịch sử: https://github.com/maihanh168-oss/tpour-trpo-experiments/tree/bfc75ca9857dd8e0a46440ce1229b22230cdefec. Số liệu 20 bước/30 query chỉ thuộc giai đoạn này.',
      '[7] Zhizhong Li, Derek Hoiem. Learning without Forgetting. https://arxiv.org/abs/1606.09282. Nền tảng bổ trợ về giữ năng lực cũ; không phải nguồn thay thế TPOUR.',
      '[8] Yingrui Yang, Yifan Qiao, Shanxiu He, Tao Yang. Weighted KL-Divergence for Document Ranking Model Refinement. https://arxiv.org/abs/2406.05977. Tài liệu liên quan về ranking distillation; biến thể nhóm không dùng weighted-KL của công trình này.',
      '[9] Kho nhóm và bằng chứng chạy: https://github.com/lequangduccode/tpour-trpo-experiments. Train retention source 219a21a, live inference source 0d2e084; results/anchored_kaggle/. SciFact/qrels từ bản phân phối BEIR, hash trong data_manifest/metrics/scifact_download.json.'
    ]
    for t in refs:doc.add_paragraph(t)
    # Correct numbering and render-sensitive formatting throughout imported content.
    for p in doc.paragraphs:
        if p.style.name=='Heading 1' and p.text!='MỤC LỤC':p.paragraph_format.page_break_before=True
        if p.text.startswith('Công thức '):
            suffix=p.text.split('. ',1)[-1] if '. ' in p.text else p.text
            if p.text.startswith('Công thức 6.'):suffix='Cập nhật reference EMA; teacher cố định không áp dụng phép cập nhật này.'
            replace(p,'Biểu thức — '+suffix)
        # Formula runs are deliberately never reconstructed from p.text.
        for run in p.runs:
            if run.text:run.text=run.text.replace('và0,74413','và 0,74413')
        if list(p._p.iter(qn('w:drawing'))):p.paragraph_format.keep_with_next=True
    old[19].style='Title'
    # Keep cover typography; normalize body spacing so single rows/lines do not spill.
    in_body=False
    for p in doc.paragraphs:
        if p.text=='MỤC LỤC':in_body=True
        if not in_body:continue
        if p.style.name=='Normal' and not list(p._p.iter(qn('m:oMath'))):
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.line_spacing=1.15
            for r in p.runs:
                if r.font.name!='Consolas':r.font.name='Times New Roman';r.font.size=Pt(12)
    for math in doc._element.iter(qn('m:oMath')):
        for r in math.iter(qn('m:r')):
            wp=r.find(qn('w:rPr'))
            if wp is None:wp=OxmlElement('w:rPr');r.insert(0,wp)
            fonts=wp.find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');wp.append(fonts)
            for a in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+a),'Cambria Math')
            for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+a),None)
    for table_index,table in enumerate(doc.tables):
        if table_index==0:continue  # The actual cover roster is preserved byte-for-byte.
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'999999');borders.append(e)
        pr=table._tbl.tblPr
        for e in list(pr):
            if e.tag==qn('w:tblBorders'):pr.remove(e)
        pr.append(borders)
        for i,row in enumerate(table.rows):
            pr=row._tr.get_or_add_trPr()
            if i==0 and pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
            if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5 if len(table.columns)>3 else 11)
                    if (len(table.rows)<=5 or table_index==3) and i<len(table.rows)-1:p.paragraph_format.keep_with_next=True
    used={e.get(qn('r:embed')) for e in doc._element.iter() if e.get(qn('r:embed'))}
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
    # Drawing IDs must remain unique after copying across documents.
    for i,e in enumerate(doc._element.iter(qn('wp:docPr')),1):e.set('id',str(i))
    for e in doc._element.iter(qn('w:instrText')):
        if e.text and e.text.strip().startswith('TOC '):e.text=' TOC \\o "1-2" \\h \\z '
    doc.sections[0].different_first_page_header_footer=True
    for part in [doc.sections[0].first_page_header,doc.sections[0].first_page_footer]:
        part.is_linked_to_previous=False
        for paragraph in part.paragraphs:paragraph.clear()
    doc.core_properties.title='TPOUR/TRPO — mẫu PTIT, công thức và kết quả Kaggle đã rà soát'
    doc.core_properties.comments='Dựng trực tiếp từ file PTIT được chỉ định; giữ cover/thành viên, không dùng file Downloads khác làm mẫu.'
    OUT.mkdir(parents=True,exist_ok=True);dest=OUT/(NAME+'.docx');doc.save(dest)
    assert digest(SOURCE)==before
    inventory=[]
    for i,p in enumerate(doc.paragraphs):
        for e in p._p.iter(qn('m:oMath')):inventory.append({'paragraph':i,'math_text':''.join(t.text or '' for t in e.iter(qn('m:t')))})
    assert all(any(key in e['math_text'] for e in inventory) for key in ['θmid','τt','LTRPO','LCE','MAP@k','MRR','Pteacher,i,j'])
    manifest={'template':str(SOURCE.relative_to(ROOT)),'template_sha256':before,'template_unchanged':True,
      'updated_evidence_source':str(UPDATE.relative_to(ROOT)),'output':str(dest.relative_to(ROOT)),
      'output_sha256_before_word_render':digest(dest),'native_math_count':len(inventory),
      'equations':inventory,'retained_cover_roster':True,'figures':'Template cover decoration retained; experimental figures generated by code.',
      'main_source':'Selected TPOUR PDF; unrelated newer report template not used.',
      'scope':'Historical partial results distinguished from full SciFact; no full temporal reproduction claim.'}
    (OUT/'report_audit.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in manifest.items() if k!='equations'},ensure_ascii=False,indent=2))

if __name__=='__main__':build()
